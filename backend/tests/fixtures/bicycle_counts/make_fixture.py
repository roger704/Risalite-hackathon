"""Generate the small synthetic bicycle-counter workbook used by offline tests."""

import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color

OUTPUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("counts.xlsx")

STATIONS = (
    ("12-PA-SCH", "Schwedter Steg", 52.535, 13.402, date(2012, 1, 1), 1),
    ("02-MI-JAN-N", "Jannowitzbrücke Nord", 52.513900, 13.418000, date(2012, 1, 1), 2),
    ("02-MI-JAN-S", "Jannowitzbrücke Süd", 52.513910, 13.418010, date(2012, 1, 1), 3),
    ("17-SK-BRE-O", "Breite Straße Ost", 52.500, 13.360, date(2012, 1, 1), 4),
    ("04-MI-NO", "Nordufer", 52.513925, 13.418035, date(2023, 10, 30), 5),
)

ORANGE = PatternFill(fill_type="solid", fgColor="FFFFC000")
YELLOW = PatternFill(fill_type="solid", fgColor="FFFFEB9C")
BLUE = PatternFill(fill_type="solid", fgColor=Color(type="theme", theme=5, tint=0.4))
DST_GAP = PatternFill(fill_type="solid", fgColor=Color(type="theme", theme=7, tint=0.6))
EMPTY_HEADER = PatternFill(fill_type="solid", fgColor="FFFFFFFF")


def _hours(day: date, omit: set[int] | None = None):
    for hour in range(24):
        if omit and hour in omit:
            continue
        yield datetime.combine(day, datetime.min.time()) + timedelta(hours=hour)


def _add_year_sheet(book, year, days, *, conflicting=False):
    sheet = book.create_sheet(f"Jahresdatei {year}")
    sheet.append(["Zählstelle        Inbetriebnahme"])
    for index, (station_id, _description, _lat, _lon, installed, _count) in enumerate(STATIONS):
        source_id = "17-SZ-BRE-O" if station_id == "17-SK-BRE-O" else station_id
        separator = ("\n", " ", "  ")[index % 3]
        sheet.cell(
            row=1,
            column=index + 2,
            value=f"{source_id}{separator}{installed.strftime('%d.%m.%Y')}",
        )
    empty_header_column = len(STATIONS) + 2
    sheet.cell(row=1, column=empty_header_column).fill = EMPTY_HEADER

    for day, gap_fill in days:
        values = []
        for hour_stamp in _hours(day, {2} if gap_fill == "absent" else None):
            row_values = [hour_stamp]
            for station_id, _description, _lat, _lon, installed, base_count in STATIONS:
                value = None if day < installed else base_count
                fill = None
                if gap_fill is True and hour_stamp.hour == 2:
                    fill = DST_GAP
                    value = base_count if day.year < 2020 and station_id == "12-PA-SCH" else None
                if day == date(2024, 1, 1) and hour_stamp.hour == 0 and conflicting:
                    if station_id == "17-SK-BRE-O":
                        value = 99
                row_values.append(value)
            values.append((hour_stamp, row_values))

        for stamp, row_values in values:
            row_number = sheet.max_row + 1
            for column, value in enumerate(row_values, 1):
                sheet.cell(row=row_number, column=column, value=value)
            for station_index, (station_id, *_rest) in enumerate(STATIONS, 2):
                cell = sheet.cell(row=row_number, column=station_index)
                fill = None
                if gap_fill is True and stamp.hour == 2:
                    fill = DST_GAP
                if day == date(2025, 2, 3) and station_id == "12-PA-SCH":
                    if stamp.hour == 0:
                        fill = YELLOW
                    elif stamp.hour == 1:
                        fill = ORANGE
                if day == date(2025, 2, 3) and stamp.hour == 11 and station_id == "04-MI-NO":
                    cell.value = 77
                    fill = DST_GAP
                if (
                    day.year == 2025
                    and station_id == "12-PA-SCH"
                    and day.weekday() < 5
                    and stamp.hour in {7, 8, 9}
                ):
                    weekday_index = sum(
                        current.weekday() < 5
                        for current in (
                            date(2025, 2, 1) + timedelta(days=offset)
                            for offset in range((day - date(2025, 2, 1)).days)
                        )
                    )
                    slot_index = weekday_index * 3 + stamp.hour - 7
                    if slot_index < 7:
                        fill = YELLOW
                if day == date(2025, 2, 3) and stamp.hour == 2 and station_id == "02-MI-JAN-N":
                    cell.value = 5
                    fill = BLUE
                if fill is not None:
                    cell.fill = fill
            # The empty-header column remains empty even on populated rows.
            sheet.cell(row=sheet.max_row, column=empty_header_column).value = None

    sheet.append([None] * sheet.max_column)
    return sheet


def create_workbook(path: Path = OUTPUT) -> None:
    book = Workbook()
    info = book.active
    info.title = "Hinweise"
    info["A1"] = "Synthetic bicycle counter workbook for connector tests."

    legend = book.create_sheet("Legende")
    for coordinate, fill in (("B2", ORANGE), ("B3", YELLOW), ("B4", BLUE)):
        legend[coordinate].fill = fill

    stations_sheet = book.create_sheet("Standortdaten")
    stations_sheet.append(
        [
            "Zählstelle",
            "Beschreibung - Fahrtrichtung",
            "Breitengrad",
            "Längengrad",
            "Installationsdatum",
        ]
    )
    for station_id, description, latitude, longitude, installed, _count in STATIONS:
        stations_sheet.append(
            [
                station_id,
                description,
                latitude,
                longitude,
                datetime.combine(installed, datetime.min.time()),
            ]
        )

    _add_year_sheet(
        book,
        2019,
        [(date(2019, 3, 31), "unmarked")],
    )
    _add_year_sheet(
        book,
        2023,
        [
            (date(2023, 3, 26), True),
            (date(2023, 10, 29), False),
            (date(2023, 10, 30), False),
            (date(2024, 1, 1), False),
        ],
    )
    _add_year_sheet(
        book,
        2024,
        [
            (date(2024, 1, 1), False),
            (date(2024, 3, 31), "absent"),
        ],
        conflicting=True,
    )
    _add_year_sheet(
        book,
        2025,
        [(date(2025, 2, 1) + timedelta(days=offset), False) for offset in range(14)],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


if __name__ == "__main__":
    create_workbook()
