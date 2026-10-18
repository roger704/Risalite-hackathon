"""Opt-in smoke test against the official Berlin bicycle-counter workbook."""

import json
import os
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.connectors.bicycle_counts import _color_signature, _matches_color
from app.domain import DatasetConfig
from app.infrastructure.http import HTTPFetcher

CONFIG = Path(__file__).parents[2] / "config/datasets/berlin_bicycle_counts.json"


@pytest.mark.live
async def test_bicycle_counts_live_workbook_signature_sheets_and_legend(tmp_path):
    if os.environ.get("RUN_LIVE_TESTS") != "1":
        pytest.skip("set RUN_LIVE_TESTS=1 to check the official workbook")
    values = json.loads(CONFIG.read_text())
    values["allowed_hosts"] = tuple(values["allowed_hosts"])
    config = DatasetConfig(**values)
    asset = await HTTPFetcher().download(config, tmp_path / "official.xlsx")
    path = Path(asset.path)

    with path.open("rb") as handle:
        assert handle.read(2) == b"PK"
    with ZipFile(path) as archive:
        assert sum(item.file_size for item in archive.infolist()) <= 400_000_000
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        assert {"Standortdaten", "Legende"}.issubset(book.sheetnames)
        for coordinate in ("B2", "B3", "B4"):
            assert _matches_color(
                _color_signature(book["Legende"][coordinate]),
                config.parser_options["legend_fills"][coordinate],
            )
    finally:
        book.close()
