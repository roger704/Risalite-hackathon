"""Build small synthetic household workbooks for offline connector tests."""

import json
from pathlib import Path

from openpyxl import Workbook

CONFIG_PATH = Path(__file__).parents[1] / "config/datasets/berlin_households.json"
DEFAULT_CODES = ("01011101", "01011102", "01011103")


def households_config(expected_area_count: int = 3):
    from app.domain import DatasetConfig

    values = json.loads(CONFIG_PATH.read_text())
    values["parser_options"]["expected_area_count"] = expected_area_count
    return DatasetConfig(**values)


def _values(index: int) -> dict[str, dict[str, int | float]]:
    households = (100, 50, 0)[index]
    single = (60, 30, 0)[index]
    multi = households - single
    residents = households * 1.5
    married = multi // 2
    married_no_children = married // 2
    married_with_children = married - married_no_children
    cohabiting = multi // 4
    cohabiting_no_children = cohabiting // 2
    cohabiting_with_children = cohabiting - cohabiting_no_children
    single_parent = multi // 8
    other = multi - married - cohabiting - single_parent
    children = married_with_children + cohabiting_with_children + single_parent
    one_child = children // 4
    two_children = children // 2
    three_plus = children - one_child - two_children
    first_age = single // 4
    second_age = single // 3
    third_age = single - first_age - second_age
    average = residents / households if households else 0.0
    published_share = 100 * single_parent / children if children else 0.0
    return {
        "T1.2": {
            "B": households,
            "C": single,
            "D": multi,
            "E": multi // 2,
            "F": multi // 4,
            "G": multi - multi // 2 - multi // 4,
            "H": average,
        },
        "T1.1": {"B": residents, "C": single, "H": average},
        "T2.2": {
            "B": households,
            "C": single,
            "D": multi,
            "E": married,
            "F": married_no_children,
            "G": married_with_children,
            "H": cohabiting,
            "I": cohabiting_no_children,
            "J": cohabiting_with_children,
            "K": single_parent,
            "L": other,
        },
        "T4": {
            "B": households,
            "C": households - children,
            "D": children,
            "E": one_child,
            "F": two_children,
            "G": three_plus,
            "H": single_parent,
            "I": published_share,
        },
        "T3": {
            "B": single,
            "C": first_age,
            "D": second_age,
            "E": third_age,
            "F": single // 2,
        },
    }


def build_workbook(tmp_path, codes=DEFAULT_CODES, edits=None):
    config = households_config(len(codes))
    options = config.parser_options
    workbook = Workbook()
    workbook.remove(workbook.active)

    directory = options["code_directory"]
    sheet = workbook.create_sheet(directory["worksheet"])
    for coordinate, value in directory["expected_cells"].items():
        sheet[coordinate] = value
    first_directory_row = directory["first_data_row"]
    for index, code in enumerate(codes):
        sheet[f"{directory['code_column']}{first_directory_row + index}"] = code

    data_by_code = {code: _values(index) for index, code in enumerate(codes)}
    for table in options["tables"]:
        sheet_name = table["worksheet"]
        sheet = workbook.create_sheet(sheet_name)
        for coordinate, value in table["expected_cells"].items():
            sheet[coordinate] = value
        first = table["first_data_row"]
        for index, code in enumerate(codes):
            row_number = first + index
            sheet[f"A{row_number}"] = code
            for column, value in data_by_code[code][sheet_name].items():
                sheet[f"{column}{row_number}"] = value

        district_row = first + len(codes)
        sheet[f"A{district_row}"] = None
        sheet[f"B{district_row}"] = "Mitte"
        aggregate_row = district_row + 1
        berlin_row = aggregate_row + 1
        sheet[f"A{aggregate_row}"] = "Zusammen"
        sheet[f"A{berlin_row}"] = options["total_label"]
        for column in table["check_columns"]:
            area_values = [data_by_code[code][sheet_name].get(column, 0) for code in codes]
            aggregate = sum(area_values)
            if sheet_name in {"T1.2", "T1.1"} and column == "H":
                denominator = sum(data_by_code[code]["T1.2"]["B"] for code in codes)
                aggregate = (
                    sum(data_by_code[code]["T1.1"]["B"] for code in codes) / denominator
                    if denominator
                    else 0.0
                )
            if sheet_name == "T4" and column == "I":
                children = sum(data_by_code[code]["T4"]["D"] for code in codes)
                aggregate = (
                    100 * sum(data_by_code[code]["T4"]["H"] for code in codes) / children
                    if children
                    else 0.0
                )
            sheet[f"{column}{aggregate_row}"] = aggregate
            sheet[f"{column}{berlin_row}"] = aggregate

    for sheet_name, changes in (edits or {}).items():
        sheet = workbook[sheet_name]
        for coordinate, value in changes.items():
            sheet[coordinate] = value
    path = Path(tmp_path) / "households.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def area_row(table: dict, index: int) -> int:
    return table["first_data_row"] + index


def berlin_row(table: dict, area_count: int) -> int:
    return table["first_data_row"] + area_count + 2
