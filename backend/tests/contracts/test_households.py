"""Offline connector contracts for the synthetic Berlin household workbook."""

import hashlib
import json
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.connectors.households import HouseholdsConnector
from app.domain import (
    Area,
    AreaObservation,
    ConfigError,
    DatasetConfig,
    IncompatibleSnapshotError,
    RawAsset,
    RawBundle,
    SchemaError,
)
from tests.households_workbook import (
    DEFAULT_CODES,
    area_row,
    berlin_row,
    build_workbook,
    households_config,
)

BACKEND = Path(__file__).parents[2]
POPULATION_CONFIG = BACKEND / "config/datasets/berlin_population.json"


def raw(path: Path, config: DatasetConfig | None = None) -> RawBundle:
    content = path.read_bytes()
    config = config or households_config()
    return RawBundle(
        assets=(
            RawAsset(
                path=str(path),
                sha256=hashlib.sha256(content).hexdigest(),
                size_bytes=len(content),
                resource_url=config.resource_url,
            ),
        ),
        snapshot_id="test-households",
        request_fingerprint="test",
        retrieved_at=datetime.now(timezone.utc),
    )


def test_synthetic_workbook_emits_reviewed_metrics_and_derived_ratio(tmp_path):
    config = households_config()
    connector = HouseholdsConnector()
    batch = connector.normalize(raw(build_workbook(tmp_path), config), config)
    assert connector.validate(batch).valid
    assert len(batch.records) == 27 * 3
    assert len({record.record_id for record in batch.records}) == len(batch.records)
    assert {record.reference_period for record in batch.records} == {"2025-12-31"}
    assert {record.hierarchy for record in batch.records} == {"planungsraum"}
    assert {record.boundary_edition for record in batch.records} == {"2021"}
    assert {"households", "persons", "persons_per_household", "percent"} == {
        record.unit for record in batch.records
    }
    records = {(record.area_code, record.metric_code): record for record in batch.records}
    average = records[DEFAULT_CODES[0], "average_household_size"]
    assert average.value == 1.5
    assert isinstance(average.value, float)
    share = records[DEFAULT_CODES[0], "households_with_children_share"]
    assert share.value == 20
    assert share.quality_flags == (
        "derived_ratio",
        "numerator:households_with_children",
        "denominator:households_total",
    )
    zero_share = records[DEFAULT_CODES[2], "households_with_children_share"]
    assert zero_share.value is None
    assert zero_share.value_status == "not_applicable"
    assert batch.warnings[-1] == "Code directory: Stand  Dezember 2024"


@pytest.mark.parametrize(
    ("problem", "message"),
    [
        ("header", "T1.2!B3"),
        ("worksheet", "T4"),
        ("html", "HTML"),
        ("directory", "Schlüssel!A2"),
    ],
)
def test_workbook_schema_and_html_changes_are_rejected(tmp_path, problem, message):
    config = households_config()
    path = build_workbook(tmp_path)
    if problem == "html":
        path.write_text("<!DOCTYPE html><html>provider error</html>")
    elif problem == "header":
        workbook = load_workbook(path)
        workbook["T1.2"]["B3"] = "Changed header"
        workbook.save(path)
    elif problem == "worksheet":
        workbook = load_workbook(path)
        workbook.remove(workbook["T4"])
        workbook.save(path)
    else:
        workbook = load_workbook(path)
        workbook["Schlüssel"]["A2"] = "Changed directory edition"
        workbook.save(path)
    with pytest.raises(SchemaError, match=message):
        HouseholdsConnector().normalize(raw(path, config), config)


def test_sentinels_preserve_suppression_zero_and_derived_missingness(tmp_path):
    config = households_config()
    table = next(t for t in config.parser_options["tables"] if t["worksheet"] == "T3")
    path = build_workbook(
        tmp_path,
        edits={
            "T4": {"D10": "•"},
            "T3": {
                f"F{area_row(table, 0)}": "–",
                f"F{berlin_row(table, len(DEFAULT_CODES))}": 15,
            },
        },
    )
    connector = HouseholdsConnector()
    batch = connector.normalize(raw(path, config), config)
    assert connector.validate(batch).valid
    records = {(record.area_code, record.metric_code): record for record in batch.records}
    suppressed = records[DEFAULT_CODES[0], "households_with_children"]
    assert suppressed.value is None
    assert suppressed.value_status == "suppressed"
    assert suppressed.quality_flags == ("source_marker:•",)
    assert records[DEFAULT_CODES[0], "households_with_children_share"].value is None
    assert records[DEFAULT_CODES[0], "households_with_children_share"].value_status == "missing"
    zero = records[DEFAULT_CODES[0], "single_person_households_female"]
    assert zero.value == 0
    assert zero.value_status == "observed"
    assert zero.quality_flags == ("source_marker:–",)
    assert any("T4!D: total check skipped" in warning for warning in batch.warnings)


def test_unknown_values_reject_the_area_and_fail_quality(tmp_path):
    config = households_config()
    path = build_workbook(tmp_path, edits={"T4": {"D10": "12p"}})
    connector = HouseholdsConnector()
    batch = connector.normalize(raw(path, config), config)
    quality = connector.validate(batch)
    assert not quality.valid
    assert batch.rejected_rows == 1
    assert not any(record.area_code == DEFAULT_CODES[0] for record in batch.records)
    assert any("unknown population value '12p'" in error for error in quality.errors)


@pytest.mark.parametrize("bad_code", [1011101, "1011101"])
def test_area_codes_must_remain_eight_digit_strings(tmp_path, bad_code):
    config = households_config()
    path = build_workbook(tmp_path, edits={"T1.2": {"A10": bad_code}})
    batch = HouseholdsConnector().normalize(raw(path, config), config)
    quality = HouseholdsConnector().validate(batch)
    assert not quality.valid
    assert any("invalid area code" in rejection for rejection in batch.rejections)
    assert any(DEFAULT_CODES[0] in record.area_code for record in batch.records) is False


def test_duplicate_and_missing_area_rows_are_rejected(tmp_path):
    config = households_config()
    duplicate_path = build_workbook(tmp_path, edits={"T1.2": {"A11": DEFAULT_CODES[0]}})
    duplicate = HouseholdsConnector().normalize(raw(duplicate_path, config), config)
    assert not HouseholdsConnector().validate(duplicate).valid
    assert any("duplicate area code" in item for item in duplicate.rejections)

    missing_path = build_workbook(
        tmp_path,
        edits={
            "T4": {
                "A12": None,
                "B12": None,
                "C12": None,
                "D12": None,
                "E12": None,
                "F12": None,
                "G12": None,
                "H12": None,
                "I12": None,
            }
        },
    )
    missing = HouseholdsConnector().normalize(raw(missing_path, config), config)
    assert not HouseholdsConnector().validate(missing).valid
    assert any("code directory mismatch" in item for item in missing.rejections)


def test_duplicate_metric_mapping_is_a_configuration_error(tmp_path):
    config = households_config()
    options = config.parser_options
    options["tables"][1]["metrics"][0]["metric_code"] = "households_total"
    path = build_workbook(tmp_path)
    with pytest.raises(ConfigError, match="Duplicate household metric mapping"):
        HouseholdsConnector().normalize(raw(path, config), config)


@pytest.mark.parametrize("issue", ["column", "reference", "age_band"])
def test_invalid_metric_columns_references_and_age_bands_are_config_errors(tmp_path, issue):
    config = households_config()
    options = config.parser_options
    if issue == "column":
        options["tables"][0]["metrics"][0]["column"] = "Z"
    elif issue == "reference":
        options["sum_checks"][0][0] = "T1.2!Z"
    else:
        options["tables"][-1]["metrics"][2]["age_min"] = 60
    path = build_workbook(tmp_path)
    with pytest.raises(ConfigError):
        HouseholdsConnector().normalize(raw(path, config), config)


def test_area_consistency_and_berlin_totals_are_enforced(tmp_path):
    config = households_config()
    table = next(t for t in config.parser_options["tables"] if t["worksheet"] == "T1.2")
    path = build_workbook(
        tmp_path,
        edits={"T1.2": {f"E{area_row(table, 0)}": 21}},
    )
    batch = HouseholdsConnector().normalize(raw(path, config), config)
    assert not HouseholdsConnector().validate(batch).valid
    assert any("sum check T1.2!D" in rejection for rejection in batch.rejections)

    total_row = berlin_row(table, 3)
    path = build_workbook(
        tmp_path,
        edits={"T1.2": {f"B{total_row}": 151}},
    )
    batch = HouseholdsConnector().normalize(raw(path, config), config)
    assert not HouseholdsConnector().validate(batch).valid
    assert any(
        "PLR sum" in rejection and "Berlin total" in rejection for rejection in batch.rejections
    )


def _areas(codes=DEFAULT_CODES, edition="2021"):
    return [
        Area(
            source_id="berlin_lor",
            snapshot_id="lor-snapshot",
            source_record_id=code,
            record_id=f"planungsraum:{code}",
            geometry={
                "type": "Polygon",
                "coordinates": [[[13.0, 52.5], [13.1, 52.5], [13.1, 52.6], [13.0, 52.5]]],
            },
            area_code=code,
            name=code,
            hierarchy="planungsraum",
            boundary_edition=edition,
        )
        for code in codes
    ]


def test_parent_join_requires_exact_area_coverage_and_boundary_edition(tmp_path):
    config = households_config()
    batch = HouseholdsConnector().normalize(raw(build_workbook(tmp_path), config), config)
    connector = HouseholdsConnector()
    joined = connector.check_parent(config, batch.records, _areas())
    assert joined["area_count"] == 3
    assert joined["metric_count"] == 27

    missing = tuple(record for record in batch.records if record.area_code != DEFAULT_CODES[-1])
    with pytest.raises(IncompatibleSnapshotError, match="1 missing areas"):
        connector.check_parent(config, missing, _areas())

    extra = tuple(
        replace(
            record,
            area_code="01019999",
            record_id=record.record_id.replace(DEFAULT_CODES[-1], "01019999"),
        )
        if record.area_code == DEFAULT_CODES[-1]
        else record
        for record in batch.records
    )
    with pytest.raises(IncompatibleSnapshotError, match="1 unmatched areas"):
        connector.check_parent(config, extra, _areas())

    with pytest.raises(IncompatibleSnapshotError, match="boundary editions differ"):
        connector.check_parent(config, batch.records, _areas(edition="old"))

    mismatched_record = replace(batch.records[0], boundary_edition="old")
    with pytest.raises(IncompatibleSnapshotError, match="boundary editions differ"):
        connector.check_parent(config, (mismatched_record, *batch.records[1:]), _areas())


def test_household_metrics_are_distinct_from_population_and_keep_their_reference_date(
    tmp_path,
):
    config = households_config()
    batch = HouseholdsConnector().normalize(raw(build_workbook(tmp_path), config), config)
    population = DatasetConfig(**json.loads(POPULATION_CONFIG.read_text()))
    population_codes = {metric["metric_code"] for metric in population.parser_options["metrics"]}
    household_codes = {record.metric_code for record in batch.records}
    assert household_codes.isdisjoint(population_codes)
    count_metrics = [record for record in batch.records if record.unit in {"households", "persons"}]
    assert all(
        record.unit == "households" or record.metric_code == "residents_in_households_total"
        for record in count_metrics
    )
    assert config.reference_period != population.reference_period
    assert {record.reference_period for record in batch.records} == {config.reference_period}


def test_validate_rejects_duplicate_observation_keys(tmp_path):
    config = households_config()
    record = AreaObservation(
        source_id=config.dataset_id,
        snapshot_id="test",
        source_record_id="T1.2!B10",
        record_id="duplicate",
        area_code=DEFAULT_CODES[0],
        hierarchy="planungsraum",
        boundary_edition=config.boundary_edition,
        metric_code="households_total",
        value=100,
        unit="households",
        reference_period=config.reference_period,
        value_status="observed",
    )
    batch = HouseholdsConnector().normalize(raw(build_workbook(tmp_path), config), config)
    duplicate_batch = replace(batch, records=(record, record))
    quality = HouseholdsConnector().validate(duplicate_batch)
    assert not quality.valid
    assert any("Duplicate household observation key" in error for error in quality.errors)
