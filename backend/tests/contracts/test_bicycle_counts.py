"""Contracts for the reviewed Berlin bicycle-counter workbook adapter."""

import hashlib
import json
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from app.connectors.bicycle_counts import BicycleCountsConnector
from app.domain import (
    DatasetConfig,
    RawAsset,
    RawBundle,
    SchemaError,
    ValidationError,
    ValueStatus,
    record_from_dict,
    record_to_dict,
)
from app.domain.models import MobilityObservation

BACKEND = Path(__file__).parents[2]
CONFIG_PATH = BACKEND / "config/datasets/berlin_bicycle_counts.json"
FIXTURE = BACKEND / "tests/fixtures/bicycle_counts/counts.xlsx"


def _config(**option_overrides):
    values = json.loads(CONFIG_PATH.read_text())
    values["allowed_hosts"] = tuple(values["allowed_hosts"])
    values["parser_options"] = {**values["parser_options"], **option_overrides}
    return DatasetConfig(**values)


def _bundle(path=FIXTURE):
    content = path.read_bytes()
    return RawBundle(
        assets=(
            RawAsset(
                path=str(path),
                sha256=hashlib.sha256(content).hexdigest(),
                size_bytes=len(content),
                resource_url="https://example.org/counts.xlsx",
            ),
        ),
        retrieved_at=datetime(2026, 10, 18, tzinfo=timezone.utc),
        snapshot_id="synthetic-bicycle-counts",
        request_fingerprint="synthetic",
    )


def _normalize(path=FIXTURE, config=None):
    connector = BicycleCountsConnector()
    batch = connector.normalize(_bundle(path), config or _config())
    return connector, batch


def _record(batch, record_id):
    return next(record for record in batch.records if record.record_id == record_id)


def _mutated_fixture(tmp_path, sheet_name, cell, *, value=None, fill=None):
    path = tmp_path / "changed.xlsx"
    book = load_workbook(FIXTURE)
    target = book[sheet_name][cell]
    if value is not None:
        target.value = value
    if fill is not None:
        target.fill = fill
    book.save(path)
    return path


def test_normalizes_station_identity_status_dst_and_duplicate_periods():
    connector, batch = _normalize()
    quality = connector.validate(batch)

    assert quality.valid
    assert batch.rejected_rows == 0
    assert len({record.location_id for record in batch.records}) == 5
    assert not any(
        record.aggregation == "day" and 2020 <= int(record.interval_start[:4]) <= 2022
        for record in batch.records
    )
    assert any(
        warning.startswith("configured years without source rows:")
        and all(year in warning for year in ("2020", "2021", "2022"))
        for warning in batch.warnings
    )
    assert _record(batch, "day:17-SK-BRE-O:2024-01-01").value == 191
    assert _record(batch, "day:17-SK-BRE-O:2024-01-01").direction == "east"
    conflict = _record(batch, "day:17-SK-BRE-O:2024-01-01")
    assert "duplicate_period_conflict" in conflict.quality_flags
    assert "duplicate_period_cells: 120" in batch.warnings
    assert any("duplicate_period_conflict_cells: 1" in warning for warning in batch.warnings)

    spring = _record(batch, "day:12-PA-SCH:2023-03-26")
    assert spring.expected_intervals == 23
    assert "dst_spring_gap" in spring.quality_flags
    early_spring = _record(batch, "day:12-PA-SCH:2019-03-31")
    assert early_spring.expected_intervals == 23
    assert "dst_spring_gap_value_dropped" in early_spring.quality_flags
    absent_spring_hour = _record(batch, "day:12-PA-SCH:2024-03-31")
    assert absent_spring_hour.expected_intervals == 23
    autumn = _record(batch, "day:12-PA-SCH:2023-10-29")
    assert autumn.expected_intervals == 24
    assert "dst_autumn_repeated_hour_undocumented" in autumn.quality_flags

    reconstructed = _record(batch, "day:02-MI-JAN-N:2025-02-03")
    assert reconstructed.value_status is ValueStatus.ESTIMATED
    assert reconstructed.value == 51
    assert reconstructed.reconstructed_intervals == 1
    incomplete = _record(batch, "day:12-PA-SCH:2025-02-03")
    assert incomplete.value is None
    assert incomplete.value_status is ValueStatus.MISSING
    assert incomplete.implausible_intervals == 1
    assert incomplete.missing_intervals == 4
    assert "incomplete_interval" in incomplete.quality_flags

    assert not any(
        record.location_id == "04-MI-NO" and record.interval_start == "2023-10-29"
        for record in batch.records
    )
    assert "coordinates_unverified" in _record(batch, "day:04-MI-NO:2025-02-03").quality_flags
    assert all(record.area_code is None for record in batch.records)
    assert all(record.mode == "bicycle" for record in batch.records)
    assert all(
        "footfall" not in record.metric_code and "pedestrian" not in record.metric_code
        for record in batch.records
    )


def test_window_means_and_insufficient_coverage_remain_explicit():
    _, batch = _normalize()
    mean = _record(batch, "window:all_hours:02-MI-JAN-N:2025")
    assert mean.value == 2.01
    assert mean.value_status is ValueStatus.ESTIMATED
    assert mean.attributes["holiday_adjustment"] == "none"
    assert mean.attributes["covered_from"] == "2025-02-01"
    assert mean.attributes["covered_to"] == "2025-02-14"

    insufficient = _record(batch, "window:weekday_07_10:12-PA-SCH:2025")
    assert insufficient.value is None
    assert insufficient.value_status is ValueStatus.MISSING
    assert insufficient.expected_intervals == 30
    assert insufficient.missing_intervals == 7
    assert "insufficient_coverage" in insufficient.quality_flags


def test_unlegended_marking_discards_numeric_value_and_flags_day_and_window():
    _, batch = _normalize()

    day = _record(batch, "day:04-MI-NO:2025-02-03")
    assert day.value is None
    assert day.value_status is ValueStatus.MISSING
    assert day.expected_intervals == 24
    assert day.observed_intervals == 23
    assert day.missing_intervals == 1
    assert "unlegended_source_marking" in day.quality_flags

    window = _record(batch, "window:all_hours:04-MI-NO:2025")
    assert window.value is not None
    assert window.missing_intervals == 1
    assert "unlegended_source_marking" in window.quality_flags
    marking_warnings = [
        warning
        for warning in batch.warnings
        if "cells carry an undocumented source marking" in warning
    ]
    assert len(marking_warnings) == 1
    assert "1 cells carry an undocumented source marking" in marking_warnings[0]
    assert "(1 numeric) were not used" in marking_warnings[0]
    assert "2025-02-03 11:00" in marking_warnings[0]


def test_mobility_observation_round_trips_through_record_registry():
    _, batch = _normalize()
    original = _record(batch, "day:17-SK-BRE-O:2024-01-01")
    decoded = record_from_dict(record_to_dict(original))
    assert isinstance(decoded, MobilityObservation)
    assert decoded == original
    assert decoded.value_status is ValueStatus.OBSERVED


def test_mobility_observation_rejects_invalid_geometry_values_and_intervals():
    _, batch = _normalize()
    record = _record(batch, "day:17-SK-BRE-O:2024-01-01")
    invalid_updates = (
        {"geometry": {"type": "LineString", "coordinates": []}},
        {"location_name": ""},
        {"value": None, "value_status": ValueStatus.OBSERVED},
        {"value": -1},
        {"value": float("nan")},
        {"expected_intervals": record.expected_intervals + 1},
        {"aggregation": "month"},
        {"aggregation": "window", "window_id": None},
    )
    for updates in invalid_updates:
        with pytest.raises(ValidationError):
            replace(record, **updates)


def test_schema_changes_are_rejected(tmp_path):
    path = _mutated_fixture(
        tmp_path,
        "Jahresdatei 2019",
        "B1",
        value="unknown-counter 01.01.2012",
    )
    with pytest.raises(SchemaError, match="Unknown station header"):
        _normalize(path)

    path = _mutated_fixture(
        tmp_path,
        "Jahresdatei 2019",
        "G2",
        value=1,
    )
    with pytest.raises(SchemaError, match="empty header"):
        _normalize(path)

    path = _mutated_fixture(
        tmp_path,
        "Standortdaten",
        "A1",
        value="Changed station header",
    )
    with pytest.raises(SchemaError, match="Standortdaten header changed"):
        _normalize(path)

    path = _mutated_fixture(
        tmp_path,
        "Jahresdatei 2019",
        "C2",
        fill=PatternFill(fill_type="solid", fgColor="FF00FF00"),
    )
    with pytest.raises(SchemaError, match="Unknown workbook fill colour"):
        _normalize(path)


def test_changed_legend_html_and_unpacked_size_are_rejected(tmp_path):
    path = _mutated_fixture(
        tmp_path,
        "Legende",
        "B2",
        fill=PatternFill(fill_type="solid", fgColor="FF00FF00"),
    )
    with pytest.raises(SchemaError, match="Legend fill changed"):
        _normalize(path)

    html = tmp_path / "not-a-workbook.xlsx"
    html.write_text("<html>blocked</html>")
    with pytest.raises(SchemaError, match="not an XLSX workbook"):
        _normalize(html)

    with pytest.raises(SchemaError, match="unpacked size budget"):
        _normalize(config=_config(max_unpacked_bytes=1))


@pytest.mark.parametrize("invalid_value", [1.5, -1, True])
def test_invalid_numeric_values_are_rejected_rows(tmp_path, invalid_value):
    path = _mutated_fixture(
        tmp_path,
        "Jahresdatei 2019",
        "B2",
        value=invalid_value,
    )
    connector, batch = _normalize(path)
    report = connector.validate(batch)
    assert batch.rejected_rows == 1
    assert not report.valid
    assert any("count must be a nonnegative integer" in item for item in report.errors)
