"""Offline checks for semantics that would materially change location rankings."""

import hashlib
import json
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.connectors.population import PopulationConnector
from app.domain import ConfigError, DatasetConfig, RawAsset, RawBundle, SchemaError

FIXTURES = Path(__file__).parents[1] / "fixtures" / "population"
CONFIG = Path(__file__).parents[2] / "config" / "datasets" / "berlin_population.json"


def config():
    return DatasetConfig(**json.loads(CONFIG.read_text()))


def raw(path):
    content = path.read_bytes()
    return RawBundle(
        assets=(
            RawAsset(
                path=str(path),
                sha256=hashlib.sha256(content).hexdigest(),
                size_bytes=len(content),
                resource_url=config().resource_url,
            ),
        ),
        snapshot_id="test-population",
        request_fingerprint="test",
        retrieved_at=datetime.now(timezone.utc),
    )


def csv_config():
    c = config()
    return replace(
        c,
        parser_options={
            "format": "csv",
            "encoding": "utf-8",
            "delimiter": ";",
            "area_code_columns": ["LOR"],
            "area_code_widths": [8],
            "hierarchy": "planungsraum",
            "expected_area_count": 3,
            "sentinels": {"•": "suppressed", "": "missing"},
            "metrics": [
                {"column": "Total", "metric_code": "population_total", "unit": "persons"},
                {
                    "column": "Under6",
                    "metric_code": "population_age_band",
                    "unit": "persons",
                    "age_min": 0,
                    "age_max_exclusive": 6,
                },
                {
                    "column": "6to15",
                    "metric_code": "population_age_band",
                    "unit": "persons",
                    "age_min": 6,
                    "age_max_exclusive": 15,
                },
                {
                    "column": "15plus",
                    "metric_code": "population_age_band",
                    "unit": "persons",
                    "age_min": 15,
                },
            ],
            "check_age_total": True,
        },
    )


def workbook(tmp_path):
    c = config()
    w = Workbook()
    s = w.active
    s.title = "T2"
    for coordinate, value in c.parser_options["expected_cells"].items():
        s[coordinate] = value
    for col, value in zip(
        "ABCDEFGHIJKLM", ["01", "10", "01", 1, 100, 10, 20, 10, 10, 10, 10, 10, 20]
    ):
        s[f"{col}7"] = value
    s["D8"] = "Zusammen"
    s["E8"] = 100
    path = tmp_path / "population.xlsx"
    w.save(path)
    return path, replace(c, parser_options={**c.parser_options, "expected_area_count": 1})


def test_csv_preserves_zero_suppression_and_nonoverlapping_cohorts():
    connector = PopulationConnector()
    batch = connector.normalize(raw(FIXTURES / "wide.csv"), csv_config())
    assert connector.validate(batch).valid
    assert len(batch.records) == 12
    assert batch.records[0].area_code == "01100101"
    assert batch.records[0].boundary_edition == "2021"
    assert batch.records[0].reference_period == "2026-06-30"
    assert all(o.value == 0 and o.value_status == "observed" for o in batch.records[4:8])
    assert all(o.value is None and o.value_status == "suppressed" for o in batch.records[8:])
    assert batch.records[8].quality_flags == ("source_marker:•",)
    assert batch.records[2].age_max_exclusive == 15


def test_reviewed_xlsx_header_codes_and_aggregate_exclusion(tmp_path):
    path, c = workbook(tmp_path)
    connector = PopulationConnector()
    batch = connector.normalize(raw(path), c)
    assert connector.validate(batch).valid
    assert len(batch.records) == 9
    assert batch.records[0].area_code == "01100101"
    assert batch.records[0].source_record_id == "T2!E7"


def test_xlsx_header_drift_rejected(tmp_path):
    path, c = workbook(tmp_path)
    from openpyxl import load_workbook

    w = load_workbook(path)
    w["T2"]["F4"] = "under 7"
    w.save(path)
    with pytest.raises(SchemaError, match="F4"):
        PopulationConnector().normalize(raw(path), c)


def test_html_masquerading_as_csv_is_rejected(tmp_path):
    path = tmp_path / "population.csv"
    path.write_text("<!DOCTYPE html><html>landing page</html>")
    with pytest.raises(SchemaError, match="HTML"):
        PopulationConnector().normalize(raw(path), csv_config())


def test_missing_csv_column_is_schema_drift(tmp_path):
    path = tmp_path / "changed.csv"
    path.write_text((FIXTURES / "wide.csv").read_text().replace("Under6", "Under7"))
    with pytest.raises(SchemaError, match="fields changed"):
        PopulationConnector().normalize(raw(path), csv_config())


def test_duplicate_keys_fail_validation(tmp_path):
    path = tmp_path / "duplicate.csv"
    path.write_text((FIXTURES / "wide.csv").read_text().replace("01100102", "01100101"))
    connector = PopulationConnector()
    report = connector.validate(connector.normalize(raw(path), csv_config()))
    assert not report.valid
    assert any("Duplicate" in e for e in report.errors)


def test_overlapping_age_configuration_is_rejected():
    c = csv_config()
    c.parser_options["metrics"][2]["age_min"] = 5
    with pytest.raises(ConfigError, match="Overlapping"):
        PopulationConnector().normalize(raw(FIXTURES / "wide.csv"), c)


@pytest.mark.parametrize(
    "old,new", [("01100101", "1100101"), ("100;10", "101;10"), ("100;10", "100;unknown")]
)
def test_invalid_rows_or_inconsistent_totals_cannot_publish(tmp_path, old, new):
    path = tmp_path / "invalid.csv"
    path.write_text((FIXTURES / "wide.csv").read_text().replace(old, new))
    connector = PopulationConnector()
    batch = connector.normalize(raw(path), csv_config())
    assert batch.rejected_rows == 1
    assert not connector.validate(batch).valid


def test_encoding_must_be_reviewed():
    c = csv_config()
    del c.parser_options["encoding"]
    with pytest.raises(ConfigError, match="encoding"):
        PopulationConnector().normalize(raw(FIXTURES / "wide.csv"), c)


def test_official_legend_zero_missing_and_not_applicable_are_distinct(tmp_path):
    path = tmp_path / "legend.csv"
    path.write_text("LOR;Total;Under6;6to15;15plus\n01100101;–;…;x;/\n")
    c = csv_config()
    c.parser_options.update(
        expected_area_count=1, check_age_total=False, sentinels=config().parser_options["sentinels"]
    )
    connector = PopulationConnector()
    batch = connector.normalize(raw(path), c)
    assert connector.validate(batch).valid
    assert [(o.value, o.value_status) for o in batch.records] == [
        (0, "observed"),
        (None, "missing"),
        (None, "not_applicable"),
        (None, "missing"),
    ]


def test_unexpected_area_count_cannot_publish(tmp_path):
    path = tmp_path / "truncated.csv"
    path.write_text("\n".join((FIXTURES / "wide.csv").read_text().splitlines()[:2]) + "\n")
    connector = PopulationConnector()
    batch = connector.normalize(raw(path), csv_config())
    assert not connector.validate(batch).valid
    assert any("Expected 3 area rows" in e for e in batch.rejections)


def test_unrecognized_extra_csv_values_cannot_be_ignored(tmp_path):
    path = tmp_path / "extra.csv"
    path.write_text(
        (FIXTURES / "wide.csv").read_text().replace("100;10;20;70", "100;10;20;70;unexpected")
    )
    with pytest.raises(SchemaError, match="row width"):
        PopulationConnector().normalize(raw(path), csv_config())
