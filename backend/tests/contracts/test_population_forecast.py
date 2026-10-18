"""Synthetic workbook contracts and forecast snapshot integration."""

import asyncio
import hashlib
import json
import shutil
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.main import create_app
from app.application.service import DataService
from app.connectors.lor import LORConnector
from app.connectors.population_forecast import PopulationForecastConnector, expected_change
from app.domain import (
    Area,
    ConfigError,
    DatasetConfig,
    FetchRequest,
    ForecastObservation,
    IncompatibleSnapshotError,
    RawAsset,
    RawBundle,
    SchemaError,
    ValidationError,
    record_from_dict,
    record_to_dict,
)

BACKEND = Path(__file__).parents[2]
CONFIG = BACKEND / "config/datasets/berlin_population_forecast.json"
SMALL_URL = (
    "https://www.berlin.de/sen/stadt/_assets/stadtdaten/stadtwissen/"
    "bevoelkerungsprognose-2024-2040/"
    "kleinraeumige_ergebnistabellen_bevoelkerungsprognose_2024-2040-_1_.xlsx?ts=1780323150"
)
OVERVIEW_URL = (
    "https://www.berlin.de/sen/stadt/_assets/stadtdaten/stadtwissen/"
    "bevoelkerungsprognose-2024-2040/"
    "uebersichtstabellen_grafiken_bevoelkerungsprognose-2024-2040-_1_.xlsx?ts=1780323147"
)
OVERVIEW_SHEET = "T1_prog. Entwicklung 2000-2040"
SMALL_SHEETS = ["00_Übersicht", "00_Berlin", "01_Test+2PGR"]
OVERVIEW_SHEETS = [
    "00_Übersicht",
    OVERVIEW_SHEET,
    "G1_Verlauf Prognose",
    "G2_Lebensbaum_2024-2040",
    "T2_Entwicklung Altersgruppen",
]


def file_pin(path: Path, role: str, url: str) -> dict:
    content = path.read_bytes()
    return {
        "role": role,
        "url": url,
        "sha256": hashlib.sha256(content).hexdigest(),
        "size_bytes": len(content),
    }


def synthetic_workbooks(directory: Path) -> tuple[Path, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    small = Workbook()
    small.remove(small.active)
    overview_sheet = small.create_sheet("00_Übersicht")
    overview_sheet["A1"] = "Synthetic overview"
    berlin = small.create_sheet("00_Berlin")
    district = small.create_sheet("01_Test+2PGR")
    for sheet, place in ((berlin, "Berlin"), (district, "district")):
        sheet["A1"] = "Bevölkerungsprognose 2024-2040  (Basis Melderegister)"
        sheet["A2"] = "mittlere Variante"
        sheet["A4"] = "Ausgangsbestand: 31.12.2024"
        sheet["A6"] = "Berlin" if place == "Berlin" else "Bezirk"
        if place == "district":
            sheet["C4"] = datetime(2024, 12, 31)
            sheet["C6"] = "01 - Synthetic District"
            sheet["U7"] = "0101 - Synthetic PGR 1"
            sheet["AM7"] = "0102 - Synthetic PGR 2"
        for age in range(99):
            sheet.cell(9 + age, 1).value = f"{age:02d} - unter {age + 1:02d}"
        sheet["A108"] = "99 und älter"
        sheet["A109"] = "Gesamt"
        for offset, year in enumerate(range(2024, 2041)):
            if place == "Berlin":
                starts = [(3, 2)]
            else:
                starts = [(3, 2), (21, 1), (39, 1)]
            for start, multiplier in starts:
                factor = 100 + offset * 0.25
                total = factor * multiplier
                for age in range(100):
                    sheet.cell(9 + age, start + offset).value = factor / 100 * multiplier
                sheet.cell(109, start + offset).value = total
        if place == "district":
            for offset in range(17):
                factor = 100 + offset * 0.25
                for start in (21, 39):
                    for age in range(100):
                        district.cell(9 + age, start + offset).value = factor / 100
                    district.cell(109, start + offset).value = factor
                for age in range(100):
                    district.cell(9 + age, 3 + offset).value = factor / 50
                district.cell(109, 3 + offset).value = factor * 2
            for offset in range(17):
                factor = 100 + offset * 0.25
                for age in range(100):
                    berlin.cell(9 + age, 3 + offset).value = factor / 50
                berlin.cell(109, 3 + offset).value = factor * 2
        for offset, year in enumerate(range(2024, 2041)):
            sheet.cell(8, 3 + offset).value = year
            if place == "district":
                sheet.cell(8, 21 + offset).value = year
                sheet.cell(8, 39 + offset).value = year
    small_path = directory / "small_area.xlsx"
    small.save(small_path)

    overview = Workbook()
    overview.remove(overview.active)
    overview.create_sheet("00_Übersicht")
    table = overview.create_sheet(OVERVIEW_SHEET)
    table["A1"] = "Prognostizierte Entwicklung der Einwohner Berlins"
    for cell, value in {
        "B3": "Mittlere Variante",
        "D3": "Obere Variante",
        "F3": "Untere Variante",
        "A4": "31.12.",
        "B4": "Einwohner ",
        "D4": "Einwohner ",
        "F4": "Einwohner ",
    }.items():
        table[cell] = value
    labels = [2000, 2005, 2010, 2015, 2020, 2024, "2025*", "2030*", "2040*"]
    for row, label in enumerate(labels, start=5):
        table.cell(row, 1).value = label
    for column, values in zip(
        (2, 4, 6),
        (
            (3897, 3915, 3964, 4006),
            (3897, 3924, 4040, 4142),
            (3897, 3911, 3942, 3894),
        ),
    ):
        for row, value in zip((10, 11, 12, 13), values):
            table.cell(row, column).value = value
    for sheet_name in OVERVIEW_SHEETS[2:]:
        overview.create_sheet(sheet_name)
    overview_path = directory / "scenario_overview.xlsx"
    overview.save(overview_path)
    return small_path, overview_path


def config_for(small_path: Path, overview_path: Path) -> DatasetConfig:
    values = json.loads(CONFIG.read_text())
    options = values["parser_options"]
    options["expected_sheet_names"] = {
        "small_area": SMALL_SHEETS,
        "scenario_overview": OVERVIEW_SHEETS,
    }
    options["expected_area_counts"] = {"bezirk": 1, "prognoseraum": 2}
    options["assets"] = [
        file_pin(small_path, "small_area", SMALL_URL),
        file_pin(overview_path, "scenario_overview", OVERVIEW_URL),
    ]
    values["allowed_hosts"] = tuple(values["allowed_hosts"])
    return DatasetConfig(**values)


def raw_bundle(config: DatasetConfig, small_path: Path, overview_path: Path) -> RawBundle:
    assets = []
    for path, url in ((small_path, SMALL_URL), (overview_path, OVERVIEW_URL)):
        content = path.read_bytes()
        assets.append(
            RawAsset(
                path=str(path),
                sha256=hashlib.sha256(content).hexdigest(),
                size_bytes=len(content),
                resource_url=url,
            )
        )
    return RawBundle(
        assets=tuple(assets),
        snapshot_id="forecast-test-snapshot",
        request_fingerprint="test",
        retrieved_at=datetime.now(timezone.utc),
    )


def forecast_config(tmp_path: Path) -> tuple[DatasetConfig, Path, Path]:
    small_path, overview_path = synthetic_workbooks(tmp_path / "workbooks")
    return config_for(small_path, overview_path), small_path, overview_path


def mutate_workbook(path: Path, sheet: str, cell: str, value) -> None:
    book = load_workbook(path)
    book[sheet][cell] = value
    book.save(path)
    book.close()


def forecast_record(**changes) -> ForecastObservation:
    values = {
        "source_id": "berlin_population_forecast",
        "snapshot_id": "test",
        "source_record_id": "00_Berlin!C109",
        "record_id": "record",
        "area_code": "0101",
        "hierarchy": "prognoseraum",
        "boundary_edition": "2021",
        "metric_code": "population_total",
        "value": 100,
        "unit": "persons",
        "reference_period": "2024-12-31",
        "value_status": "observed",
        "scenario": "middle",
        "forecast_edition": "2024-2040",
        "base_period": "2024-12-31",
    }
    values.update(changes)
    return ForecastObservation(**values)


def test_synthetic_workbooks_emit_all_geographies_scenarios_and_round_trip(tmp_path):
    config, small, overview = forecast_config(tmp_path)
    connector = PopulationForecastConnector()
    batch = connector.normalize(raw_bundle(config, small, overview), config)
    assert connector.validate(batch).valid
    assert len(batch.records) == 4 * 17 * 101 + 12
    assert {record.scenario for record in batch.records} == {"middle", "upper", "lower"}
    assert len({record.record_id for record in batch.records}) == len(batch.records)
    assert {record.hierarchy for record in batch.records} == {
        "land",
        "bezirk",
        "prognoseraum",
    }
    assert {record.area_code for record in batch.records if record.hierarchy == "prognoseraum"} == {
        "0101",
        "0102",
    }
    observed = [record for record in batch.records if record.reference_period == "2024-12-31"]
    forecasts = [record for record in batch.records if record.reference_period == "2025-12-31"]
    assert all(record.value_status == "observed" for record in observed)
    assert all(record.value_status == "forecast" for record in forecasts)
    sample = next(record for record in forecasts if record.area_code == "0101")
    assert sample.value == 100.25
    decoded = record_from_dict(record_to_dict(sample))
    assert decoded == sample
    upper = next(
        record
        for record in batch.records
        if record.scenario == "upper" and record.reference_period == "2040-12-31"
    )
    assert upper.unit == "thousand_persons"
    assert upper.quality_flags == ("rounded_to_thousands",)


@pytest.mark.parametrize(
    ("path_name", "sheet", "cell", "value", "message"),
    [
        ("small", "00_Berlin", "A1", "Changed title", "title changed"),
        ("small", "00_Berlin", "A9", "00 - unter 02", "Age label changed"),
        ("small", "01_Test+2PGR", "A2", "obere Variante", "Unexpected forecast scenario"),
        ("overview", OVERVIEW_SHEET, "B3", "Obere Variante", "scenario header changed"),
        ("small", "00_Berlin", "D8", 2026, "year headers changed"),
        ("overview", OVERVIEW_SHEET, "A12", "2030", "year label changed"),
        ("overview", OVERVIEW_SHEET, "A12", "bad period", "year label changed"),
    ],
)
def test_changed_headers_mixed_scenarios_and_bad_periods_are_rejected(
    tmp_path, path_name, sheet, cell, value, message
):
    config, small, overview = forecast_config(tmp_path)
    target = small if path_name == "small" else overview
    mutate_workbook(target, sheet, cell, value)
    config = config_for(small, overview)
    with pytest.raises(SchemaError, match=message):
        PopulationForecastConnector().normalize(raw_bundle(config, small, overview), config)


def test_single_asset_bundle_is_a_clear_config_error(tmp_path):
    config, small, overview = forecast_config(tmp_path)
    raw = raw_bundle(config, small, overview)
    with pytest.raises(ConfigError, match="both pinned workbooks"):
        PopulationForecastConnector().normalize(replace(raw, assets=raw.assets[:1]), config)


def test_missing_cell_is_rejected_and_never_emitted_as_zero(tmp_path):
    config, small, overview = forecast_config(tmp_path)
    mutate_workbook(small, "00_Berlin", "C9", None)
    config = config_for(small, overview)
    batch = PopulationForecastConnector().normalize(raw_bundle(config, small, overview), config)
    assert batch.rejected_rows == 1
    assert any("00_Berlin!C9" in reason for reason in batch.rejections)
    assert not any(record.source_record_id == "00_Berlin!C9" for record in batch.records)
    assert not PopulationForecastConnector().validate(batch).valid


def test_age_total_discrepancy_over_one_rejects_batch(tmp_path):
    config, small, overview = forecast_config(tmp_path)
    mutate_workbook(small, "01_Test+2PGR", "U9", 3)
    config = config_for(small, overview)
    with pytest.raises(ValidationError, match="differ from total"):
        PopulationForecastConnector().normalize(raw_bundle(config, small, overview), config)


def test_prognoseraum_district_discrepancy_is_a_warning(tmp_path):
    config, small, overview = forecast_config(tmp_path)
    mutate_workbook(small, "01_Test+2PGR", "U9", 6)
    mutate_workbook(small, "01_Test+2PGR", "U109", 105)
    config = config_for(small, overview)
    connector = PopulationForecastConnector()
    batch = connector.normalize(raw_bundle(config, small, overview), config)
    assert connector.validate(batch).valid
    assert any("1 differences over 1 person" in warning for warning in batch.warnings)


def test_fetch_rejects_asset_pin_mismatch(tmp_path):
    config, small, overview = forecast_config(tmp_path)

    class FakeFetcher:
        async def download(self, selected_config, destination):
            source = small if selected_config.resource_url == SMALL_URL else overview
            shutil.copyfile(source, destination)
            content = Path(destination).read_bytes()
            return RawAsset(
                path=str(destination),
                sha256="0" * 64,
                size_bytes=len(content),
                resource_url=selected_config.resource_url,
            )

    staging = tmp_path / "download"
    staging.mkdir()
    with pytest.raises(ValidationError, match="Pinned edition changed"):
        asyncio.run(
            PopulationForecastConnector(FakeFetcher()).fetch(
                FetchRequest(
                    config=config,
                    staging_dir=str(staging),
                    snapshot_id="fetch-test",
                )
            )
        )


@pytest.mark.parametrize(
    "changes",
    [
        {"scenario": "upper"},
        {"hierarchy": "bezirk", "area_code": "01"},
        {"unit": "thousand_persons"},
        {"age_min": 1, "age_max_exclusive": 2},
        {"base_period": "2023-12-31"},
    ],
)
def test_expected_change_rejects_incompatible_forecast_pairs(changes):
    base = forecast_record()
    target = forecast_record(
        record_id="target",
        value=110,
        reference_period="2025-12-31",
        value_status="forecast",
    )
    target = replace(target, **changes)
    with pytest.raises(ValidationError):
        expected_change(base, target)


def test_expected_change_returns_percentage_and_zero_base_null():
    base = forecast_record(value=100)
    target = forecast_record(
        record_id="target",
        value=110,
        reference_period="2040-12-31",
        value_status="forecast",
    )
    change = expected_change(base, target)
    assert change["absolute_change"] == 10
    assert change["percent_change"] == 10
    assert change["hierarchy"] == "prognoseraum"
    assert "not a measured planning-area or address-level value" in change["resolution_note"]
    zero_base = replace(base, value=0)
    assert expected_change(zero_base, target)["percent_change"] is None


@pytest.mark.parametrize(
    ("changes", "message"),
    [
        ({"reference_period": "2024-12-31", "value_status": "forecast"}, "value_status"),
        ({"reference_period": "2023-12-31", "base_period": "2024-12-31"}, "before base_period"),
        ({"scenario": "medium"}, "scenario"),
        ({"area_code": "1301"}, "Invalid area code"),
    ],
)
def test_forecast_observation_rejects_invalid_periods_codes_and_scenarios(changes, message):
    with pytest.raises(ValidationError, match=message):
        forecast_record(**changes)


def parent_area(area_code: str, boundary_edition: str = "2021") -> Area:
    return Area(
        source_id="berlin_lor",
        snapshot_id="parent",
        source_record_id=area_code,
        record_id=area_code,
        geometry={
            "type": "Polygon",
            "coordinates": [[[13, 52], [13.1, 52], [13.1, 52.1], [13, 52.1], [13, 52]]],
        },
        area_code=area_code,
        name=area_code,
        hierarchy="planungsraum",
        boundary_edition=boundary_edition,
    )


def joined_forecasts():
    return (
        forecast_record(area_code="0101", hierarchy="prognoseraum"),
        forecast_record(record_id="district", area_code="01", hierarchy="bezirk"),
        forecast_record(record_id="second-pgr", area_code="0102", hierarchy="prognoseraum"),
    )


def test_check_parent_requires_exact_prognoseraum_district_prefix_sets(tmp_path):
    config, _, _ = forecast_config(tmp_path)
    connector = PopulationForecastConnector()
    forecasts = joined_forecasts()
    parent = (parent_area("01010101"), parent_area("01020101"))
    coverage = connector.check_parent(config, forecasts, parent)
    assert coverage["prognoseraum_count"] == 2
    assert coverage["bezirk_count"] == 1
    assert coverage["resolution"] == "prognoseraum"
    with pytest.raises(IncompatibleSnapshotError, match="1 missing"):
        connector.check_parent(config, (forecasts[0], forecasts[1]), parent)
    with pytest.raises(IncompatibleSnapshotError, match="1 unmatched"):
        connector.check_parent(
            config,
            (*forecasts, forecast_record(record_id="extra-pgr", area_code="0103")),
            parent,
        )
    with pytest.raises(IncompatibleSnapshotError, match="missing"):
        connector.check_parent(
            config,
            forecasts,
            (parent_area("01010101", "2020"), parent_area("01020101", "2020")),
        )


class FixtureFetcher:
    def __init__(self, small_path: Path, overview_path: Path):
        self.small_path = small_path
        self.overview_path = overview_path
        self.bad_small_path: Path | None = None

    async def download(self, config: DatasetConfig, destination: Path) -> RawAsset:
        source = (
            self.bad_small_path
            if config.resource_url == SMALL_URL and self.bad_small_path is not None
            else self.small_path
            if config.resource_url == SMALL_URL
            else self.overview_path
        )
        shutil.copyfile(source, destination)
        content = Path(destination).read_bytes()
        return RawAsset(
            path=str(destination),
            sha256=hashlib.sha256(content).hexdigest(),
            size_bytes=len(content),
            resource_url=config.resource_url,
        )


async def service_fixture(tmp_path: Path):
    small, overview = synthetic_workbooks(tmp_path / "source")
    forecast_config_obj = config_for(small, overview)
    config_dir = tmp_path / "config"
    config_dir.mkdir()
    lor_values = json.loads((BACKEND / "config/datasets/berlin_lor.json").read_text())
    lor_values["source_crs"] = "EPSG:4326"
    lor_values["parser_options"].update(expected_rows=2, name_field="PLR_NAME")
    (config_dir / "berlin_lor.json").write_text(json.dumps(lor_values))
    forecast_values = json.loads(CONFIG.read_text())
    forecast_values["redistribution_status"] = "allowed_with_attribution"
    forecast_values["parser_options"]["expected_sheet_names"] = forecast_config_obj.parser_options[
        "expected_sheet_names"
    ]
    forecast_values["parser_options"]["expected_area_counts"] = {"bezirk": 1, "prognoseraum": 2}
    forecast_values["parser_options"]["assets"] = forecast_config_obj.parser_options["assets"]
    (config_dir / "berlin_population_forecast.json").write_text(json.dumps(forecast_values))
    lor_file = tmp_path / "parent.geojson"
    lor_data = json.loads((BACKEND / "tests/fixtures/lor/areas.geojson").read_text())
    lor_data["features"][1]["properties"]["PLR_ID"] = "01021102"
    lor_file.write_text(json.dumps(lor_data))
    fetcher = FixtureFetcher(small, overview)
    service = DataService(
        data_dir=tmp_path / "data",
        config_dir=config_dir,
        connectors={
            "lor": LORConnector(),
            "population_forecast": PopulationForecastConnector(fetcher),
        },
    )
    parent = await service.ingest("berlin_lor", local_file=lor_file)
    forecast = await service.ingest(
        "berlin_population_forecast", parent_snapshot_id=parent.snapshot_id
    )
    return service, fetcher, small, overview, parent, forecast


@pytest.mark.asyncio
async def test_service_ingest_api_and_failed_refresh_preserve_latest_snapshot(tmp_path):
    service, fetcher, small, overview, parent, snapshot = await service_fixture(tmp_path)
    assert snapshot.parent_snapshot_ids == (parent.snapshot_id,)
    assert len(snapshot.raw_assets) == 2
    assert snapshot.raw_assets[0].resource_url == SMALL_URL
    assert snapshot.raw_assets[1].resource_url == OVERVIEW_URL
    assert snapshot.license_id == "LicenseRef-berlin-de-unstated"
    assert snapshot.attribution.startswith("Senatsverwaltung für Stadtentwicklung")
    assert snapshot.coverage["prognoseraum_count"] == 2
    assert snapshot.coverage["join_method"].startswith("Prognoseraum/Bezirk")
    assert snapshot.quality_report.warnings

    with TestClient(create_app(service=service)) as client:
        response = client.get(
            "/api/v1/population-forecast/expected-change",
            params={
                "snapshot_id": snapshot.snapshot_id,
                "area_code": "0101",
                "target_year": 2040,
                "scenario": "middle",
            },
        )
        assert response.status_code == 200, response.text
        change = response.json()
        assert change["base_value"] == 100
        assert change["target_value"] == 104
        assert change["scenario"] == "middle"
        assert change["target_period"] == "2040-12-31"
        assert change["hierarchy"] == "prognoseraum"
        assert "not a measured planning-area or address-level value" in change["resolution_note"]
        planungsraum = client.get(
            "/api/v1/population-forecast/expected-change",
            params={
                "snapshot_id": snapshot.snapshot_id,
                "area_code": "01010101",
                "target_year": 2040,
            },
        )
        assert planungsraum.status_code == 422
        assert "Prognoseraum" in planungsraum.json()["error"]["message"]
        unknown_scenario = client.get(
            "/api/v1/population-forecast/expected-change",
            params={
                "snapshot_id": snapshot.snapshot_id,
                "area_code": "0101",
                "target_year": 2040,
                "scenario": "medium",
            },
        )
        assert unknown_scenario.status_code == 422
        unknown_year = client.get(
            "/api/v1/population-forecast/expected-change",
            params={
                "snapshot_id": snapshot.snapshot_id,
                "area_code": "0101",
                "target_year": 2041,
            },
        )
        assert unknown_year.status_code == 422
        unavailable = client.get(
            "/api/v1/population-forecast/expected-change",
            params={
                "snapshot_id": snapshot.snapshot_id,
                "area_code": "0101",
                "target_year": 2040,
                "scenario": "upper",
            },
        )
        assert unavailable.status_code == 404
        record_page = client.get(
            "/api/v1/datasets/berlin_population_forecast/records",
            params={"snapshot_id": snapshot.snapshot_id, "area_code": "0101", "limit": 1},
        )
        assert record_page.status_code == 200
        assert record_page.json()["records"][0]["scenario"] == "middle"
        assert record_page.json()["records"][0]["value_status"] in {"observed", "forecast"}

    bad_workbook = tmp_path / "broken.xlsx"
    bad_workbook.write_bytes(b"not an xlsx workbook")
    fetcher.bad_small_path = bad_workbook
    updated = service.configs["berlin_population_forecast"]
    options = dict(updated.parser_options)
    options["assets"] = [
        file_pin(bad_workbook, "small_area", SMALL_URL),
        file_pin(overview, "scenario_overview", OVERVIEW_URL),
    ]
    service.configs["berlin_population_forecast"] = replace(updated, parser_options=options)
    with pytest.raises(SchemaError, match="valid XLSX archive"):
        await service.ingest("berlin_population_forecast", parent_snapshot_id=parent.snapshot_id)
    latest = service.catalog.get("berlin_population_forecast")
    assert latest.snapshot_id == snapshot.snapshot_id
    readable = service.records(
        "berlin_population_forecast", snapshot.snapshot_id, area_code="0101", limit=1
    )
    assert readable["count"] == 1
