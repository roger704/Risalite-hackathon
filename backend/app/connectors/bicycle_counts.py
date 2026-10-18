"""Berlin's reviewed hourly bicycle-counter workbook."""

from __future__ import annotations

import math
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.connectors.base import FileConnector
from app.domain import (
    ConfigError,
    DatasetConfig,
    NormalizedBatch,
    QualityReport,
    RawBundle,
    SchemaError,
)
from app.domain.models import MobilityObservation

_YEAR_SHEET = re.compile(r"^Jahresdatei\s+(\d{4})$")
_HEADER = re.compile(r"^(\S+)\s+(\d{2}\.\d{2}\.\d{4})$")
_DIRECTIONS = {"Nord": "north", "Süd": "south", "Ost": "east", "West": "west"}
_ID_DIRECTIONS = {"N": "north", "S": "south", "O": "east", "W": "west"}
_COUNTS = ("observed", "reconstructed", "implausible")


def _normalized_space(value: object) -> str:
    return " ".join(str(value or "").split())


def _empty(value: object) -> bool:
    return value is None or value == ""


def _color_signature(cell) -> tuple[str, object, float] | None:
    fill = getattr(cell, "fill", None)
    if fill is None or fill.fill_type in {None, "none"}:
        return None
    color = fill.fgColor
    if color.type == "rgb":
        return ("rgb", str(color.rgb).upper(), 0.0)
    if color.type == "theme":
        return ("theme", color.theme, float(color.tint or 0.0))
    raise SchemaError(f"Unsupported workbook fill color type: {color.type}")


def _matches_color(signature, expected: dict) -> bool:
    if signature is None or signature[0] != expected.get("type"):
        return False
    if signature[1] != expected.get("value"):
        return False
    return expected.get("type") != "theme" or math.isclose(
        signature[2], float(expected.get("tint", 0.0)), abs_tol=0.01
    )


@dataclass(frozen=True)
class _SourceCell:
    value: int | None
    raw_value: object
    value_class: str
    sheet: str
    sheet_year: int
    row: int
    column: int
    order: tuple[int, int]

    @property
    def source_record_id(self) -> str:
        column = get_column_letter(self.column)
        return f"{self.sheet}!{column}{self.row}..{self.row}"


@dataclass
class _Totals:
    observed: int = 0
    reconstructed: int = 0
    implausible: int = 0
    value_sum: int = 0
    flags: set[str] = field(default_factory=set)
    source_order: tuple[int, int] | None = None
    source_record_id: str | None = None
    source_sheets: set[str] = field(default_factory=set)

    def add(self, cell: _SourceCell, conflict: bool = False) -> None:
        if cell.value_class == "observed":
            self.observed += 1
            self.value_sum += cell.value or 0
        elif cell.value_class == "reconstructed":
            self.reconstructed += 1
            self.value_sum += cell.value or 0
        elif cell.value_class == "implausible":
            self.implausible += 1
        elif cell.value_class == "unlegended_source_marking":
            self.flags.add("unlegended_source_marking")
        self.source_sheets.add(cell.sheet)
        if self.source_order is None or cell.order < self.source_order:
            self.source_order = cell.order
            self.source_record_id = cell.source_record_id
        if conflict:
            self.flags.add("duplicate_period_conflict")


@dataclass(frozen=True)
class _Window:
    window_id: str
    label: str
    days: str
    start_hour: int
    end_hour: int


class BicycleCountsConnector(FileConnector):
    adapter_version = "1.0.0"

    def normalize(self, raw: RawBundle, config: DatasetConfig) -> NormalizedBatch:
        if len(raw.assets) != 1:
            raise SchemaError("Bicycle counts import requires exactly one XLSX asset")
        options = config.parser_options
        years = self._years(options.get("years"))
        windows = self._windows(options.get("windows"))
        timezone_name = options.get("timezone")
        if not isinstance(timezone_name, str) or not timezone_name:
            raise ConfigError("timezone must be an IANA timezone name")
        try:
            zone = ZoneInfo(timezone_name)
        except (ZoneInfoNotFoundError, ValueError) as exc:
            raise ConfigError(f"Unknown timezone: {timezone_name}") from exc
        min_coverage = options.get("min_window_coverage")
        if (
            isinstance(min_coverage, bool)
            or not isinstance(min_coverage, (int, float))
            or not math.isfinite(min_coverage)
            or not 0 <= min_coverage <= 1
        ):
            raise ConfigError("min_window_coverage must be between 0 and 1")
        max_unpacked = options.get("max_unpacked_bytes")
        if type(max_unpacked) is not int or max_unpacked <= 0:
            raise ConfigError("max_unpacked_bytes must be a positive integer")
        reviewed_sha256 = options.get("reviewed_sha256")
        if not isinstance(reviewed_sha256, str) or not re.fullmatch(
            r"[0-9a-fA-F]{64}", reviewed_sha256
        ):
            raise ConfigError("reviewed_sha256 must be a 64-character SHA-256 digest")

        source = Path(raw.assets[0].path)
        try:
            with source.open("rb") as handle:
                prefix = handle.read(512)
            if not prefix.startswith(b"PK") or prefix.lstrip().lower().startswith(
                (b"<!doctype html", b"<html")
            ):
                raise SchemaError("Bicycle counts resource is not an XLSX workbook")
            with ZipFile(source) as archive:
                if sum(item.file_size for item in archive.infolist()) > max_unpacked:
                    raise SchemaError("Bicycle counts workbook exceeds unpacked size budget")
            book = load_workbook(source, read_only=True, data_only=True)
        except (BadZipFile, OSError, ValueError) as exc:
            raise SchemaError("Bicycle counts resource is not a readable XLSX workbook") from exc

        warnings = list(options.get("limitations", ()))
        if raw.assets[0].sha256.lower() != reviewed_sha256.lower():
            warnings.append("source file differs from reviewed edition")
        try:
            required = {"Hinweise", "Legende", "Standortdaten"}
            if not required.issubset(book.sheetnames):
                raise SchemaError(
                    f"Bicycle counts workbook is missing sheets: {sorted(required - set(book.sheetnames))}"
                )
            self._check_legend(book["Legende"], options)
            stations = self._read_stations(book["Standortdaten"])
            station_flags, coordinate_pairs = self._station_flags(stations)
            for first, second, distance in coordinate_pairs:
                warnings.append(
                    f"coordinates_unverified: {first} and {second} are {distance:.1f} m apart."
                )
            aliases = options.get("station_aliases", {})
            if not isinstance(aliases, dict):
                raise ConfigError("station_aliases must be an object")
            sheet_descriptors = []
            ignored_sheets = []
            for name in book.sheetnames:
                match = _YEAR_SHEET.fullmatch(name)
                if not match:
                    continue
                sheet_year = int(match.group(1))
                if sheet_year not in years:
                    ignored_sheets.append(name)
                else:
                    sheet_descriptors.append((name, sheet_year))
            if ignored_sheets:
                warnings.append(
                    "Ignored year sheets outside configured years: " + ", ".join(ignored_sheets)
                )
            if not sheet_descriptors:
                raise SchemaError("Bicycle counts workbook has no configured year sheets")

            headers_by_sheet = {}
            empty_header_columns = {}
            for sheet_name, sheet_year in sheet_descriptors:
                headers_by_sheet[sheet_name] = self._read_year_headers(
                    book[sheet_name], sheet_name, sheet_year, stations, aliases
                )
                header_values = next(
                    book[sheet_name].iter_rows(
                        min_row=1,
                        max_row=1,
                        max_col=book[sheet_name].max_column,
                        values_only=True,
                    )
                )
                empty_header_columns[sheet_name] = {
                    index for index, value in enumerate(header_values[1:], 2) if value is None
                }

            timestamp_counts: dict[datetime, int] = defaultdict(int)
            year_dates: dict[int, list[date]] = defaultdict(list)
            for sheet_name, _ in sheet_descriptors:
                sheet = book[sheet_name]
                for (value,) in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
                    if isinstance(value, datetime) and value.year in years:
                        timestamp_counts[value] += 1
                        year_dates[value.year].append(value.date())
            duplicate_times = {stamp for stamp, count in timestamp_counts.items() if count > 1}
            all_dates = [day for dates in year_dates.values() for day in dates]
            first_date = min(all_dates) if all_dates else None
            last_date = max(all_dates) if all_dates else None
            for station_id, station in stations.items():
                if station.get("install_date_mismatch"):
                    station_flags[station_id].add("install_date_mismatch")
                if station.get("direction_label_conflict"):
                    station_flags[station_id].add("direction_label_conflict")

            day_totals: dict[tuple[str, date], _Totals] = defaultdict(_Totals)
            window_totals: dict[tuple[str, int, str], _Totals] = defaultdict(_Totals)
            special_flags: dict[tuple[str, date], set[str]] = defaultdict(set)
            duplicate_candidates: dict[tuple[str, datetime], list[_SourceCell]] = defaultdict(list)
            rejected_rows: set[tuple[str, int]] = set()
            rejections: list[str] = []
            unlegended_cell_count = 0
            unlegended_numeric_count = 0
            unlegended_examples: set[str] = set()
            input_rows = 0
            station_first_cell: dict[str, _SourceCell] = {}

            for sheet_order, (sheet_name, sheet_year) in enumerate(sheet_descriptors):
                sheet = book[sheet_name]
                headers = headers_by_sheet[sheet_name]
                max_column = sheet.max_column
                for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=max_column), 2):
                    values = [cell.value for cell in row]
                    if all(_empty(value) for value in values):
                        continue
                    input_rows += 1
                    for column_number in empty_header_columns[sheet_name]:
                        if not _empty(values[column_number - 1]):
                            raise SchemaError(
                                f"Data exists under an empty header at {sheet_name}!{get_column_letter(column_number)}{row_number}"
                            )
                    stamp = values[0]
                    if not isinstance(stamp, datetime) or stamp.tzinfo is not None:
                        rejected_rows.add((sheet_name, row_number))
                        rejections.append(
                            f"{sheet_name}!A{row_number}: nonempty row has a non-datetime timestamp"
                        )
                        continue
                    if stamp.minute or stamp.second or stamp.microsecond:
                        rejected_rows.add((sheet_name, row_number))
                        rejections.append(
                            f"{sheet_name}!A{row_number}: timestamp is not an hourly label"
                        )
                        continue
                    if stamp.year not in years:
                        continue
                    spring_gap = self._spring_gap_day(stamp.date(), zone) and stamp.hour == 2
                    for station_id, column_number in headers.items():
                        cell = row[column_number - 1]
                        fill_class = self._cell_fill(cell, options)
                        if fill_class == "dst_gap" and not spring_gap:
                            fill_class = "unlegended_source_marking"
                            unlegended_cell_count += 1
                            if isinstance(cell.value, (int, float)) and not isinstance(
                                cell.value, bool
                            ):
                                unlegended_numeric_count += 1
                            unlegended_examples.add(stamp.strftime("%Y-%m-%d %H:%M"))
                        if fill_class == "unlegended_source_marking":
                            value = None
                        else:
                            try:
                                value = self._count_value(cell.value)
                            except ValueError as exc:
                                rejected_rows.add((sheet_name, row_number))
                                rejections.append(
                                    f"{sheet_name}!{get_column_letter(column_number)}{row_number}: {exc}"
                                )
                                continue
                        value_class = self._value_class(fill_class, value)
                        source_cell = _SourceCell(
                            value=value,
                            raw_value=(
                                None if fill_class == "unlegended_source_marking" else cell.value
                            ),
                            value_class=value_class,
                            sheet=sheet_name,
                            sheet_year=sheet_year,
                            row=row_number,
                            column=column_number,
                            order=(sheet_order, row_number),
                        )
                        station_first_cell.setdefault(station_id, source_cell)
                        if stamp in duplicate_times:
                            duplicate_candidates[(station_id, stamp)].append(source_cell)
                        else:
                            self._apply_cell(
                                source_cell,
                                station_id,
                                stamp,
                                stations,
                                windows,
                                zone,
                                day_totals,
                                window_totals,
                                special_flags,
                            )

            duplicate_cell_count = 0
            conflict_cell_count = 0
            for (station_id, stamp), candidates in duplicate_candidates.items():
                if len(candidates) > 1:
                    duplicate_cell_count += len(candidates) - 1
                    distinct = {
                        (candidate.raw_value, candidate.value_class) for candidate in candidates
                    }
                    conflict = len(distinct) > 1
                    if conflict:
                        conflict_cell_count += 1
                        preferred = [
                            candidate
                            for candidate in candidates
                            if candidate.sheet_year == stamp.year
                        ]
                        selected = min(preferred or candidates, key=lambda item: item.order)
                    else:
                        selected = min(candidates, key=lambda item: item.order)
                else:
                    selected = candidates[0]
                    conflict = False
                self._apply_cell(
                    selected,
                    station_id,
                    stamp,
                    stations,
                    windows,
                    zone,
                    day_totals,
                    window_totals,
                    special_flags,
                    conflict=conflict,
                )

            if duplicate_cell_count:
                warnings.append(f"duplicate_period_cells: {duplicate_cell_count}")
            if conflict_cell_count:
                warnings.append(
                    f"duplicate_period_conflict_cells: {conflict_cell_count}; year-matching sheet selected"
                )
            else:
                warnings.append("duplicate_period_conflict_cells: 0")

            mismatch_stations = sorted(
                station_id
                for station_id, flags in station_flags.items()
                if "install_date_mismatch" in flags
            )
            if mismatch_stations:
                warnings.append("install_date_mismatches: " + ", ".join(mismatch_stations))
            if any(self._spring_gap_day(day, zone) for day in all_dates):
                warnings.append("Spring DST-gap hours are excluded from expected intervals.")
            if any(self._autumn_change_day(day, zone) for day in all_dates):
                warnings.append("Autumn repeated-hour handling is undocumented by the source.")

            records = []
            data_years = sorted(year_dates)
            if first_date is not None and last_date is not None:
                for station_id in sorted(stations):
                    station = stations[station_id]
                    start = max(first_date, station["install_date"])
                    if start > last_date:
                        continue
                    for day in self._date_range(start, last_date):
                        if day.year not in years:
                            continue
                        totals = day_totals[(station_id, day)]
                        expected = 24 - int(self._spring_gap_day(day, zone))
                        missing = (
                            expected - totals.observed - totals.reconstructed - totals.implausible
                        )
                        if missing < 0:
                            raise SchemaError(
                                f"More hourly counts than expected for {station_id} {day}"
                            )
                        flags = (
                            set(station_flags[station_id])
                            | totals.flags
                            | special_flags[(station_id, day)]
                        )
                        if self._spring_gap_day(day, zone):
                            flags.add("dst_spring_gap")
                        if self._autumn_change_day(day, zone):
                            flags.add("dst_autumn_repeated_hour_undocumented")
                        complete = missing == 0 and totals.implausible == 0
                        value = totals.value_sum if complete else None
                        value_status = (
                            "estimated"
                            if complete and totals.reconstructed
                            else "observed"
                            if complete
                            else "missing"
                        )
                        if not complete:
                            flags.add("incomplete_interval")
                        fallback = station_first_cell.get(station_id)
                        source_id = totals.source_record_id or (
                            fallback.source_record_id if fallback else f"Standortdaten!{station_id}"
                        )
                        source_sheets = sorted(totals.source_sheets)
                        if not source_sheets and fallback:
                            source_sheets = [fallback.sheet]
                        records.append(
                            MobilityObservation(
                                source_id=config.dataset_id,
                                snapshot_id=raw.snapshot_id,
                                source_record_id=source_id,
                                record_id=f"day:{station_id}:{day.isoformat()}",
                                geometry=station["geometry"],
                                location_id=station_id,
                                location_name=station["description"],
                                mode="bicycle",
                                metric_code="bicycle_count_day",
                                unit="bicycles",
                                aggregation="day",
                                interval_start=day.isoformat(),
                                interval_end=(day + timedelta(days=1)).isoformat(),
                                timezone=timezone_name,
                                measurement_status="measured",
                                value=value,
                                value_status=value_status,
                                direction=station["direction"],
                                expected_intervals=expected,
                                observed_intervals=totals.observed,
                                reconstructed_intervals=totals.reconstructed,
                                implausible_intervals=totals.implausible,
                                missing_intervals=missing,
                                quality_flags=tuple(sorted(flags)),
                                attributes={
                                    "install_date": station["install_date"].isoformat(),
                                    "source_sheets": ", ".join(source_sheets),
                                },
                            )
                        )

                    for year in data_years:
                        year_first = max(
                            min(year_dates[year]), date(year, 1, 1), station["install_date"]
                        )
                        year_last = min(max(year_dates[year]), date(year, 12, 31))
                        if year_first > year_last:
                            continue
                        for window in windows:
                            totals = window_totals[(station_id, year, window.window_id)]
                            expected = self._window_expected(year_first, year_last, window, zone)
                            missing = (
                                expected
                                - totals.observed
                                - totals.reconstructed
                                - totals.implausible
                            )
                            if missing < 0:
                                raise SchemaError(
                                    f"More hourly counts than expected for {station_id} {year} {window.window_id}"
                                )
                            usable = totals.observed + totals.reconstructed
                            coverage = usable / expected if expected else 0.0
                            flags = set(station_flags[station_id]) | totals.flags
                            for day in self._date_range(year_first, year_last):
                                if self._window_day_matches(window, day):
                                    flags.update(special_flags[(station_id, day)])
                                    if self._spring_gap_day(day, zone):
                                        flags.add("dst_spring_gap")
                                    if self._autumn_change_day(day, zone):
                                        flags.add("dst_autumn_repeated_hour_undocumented")
                            valid = bool(expected) and coverage >= min_coverage and usable > 0
                            if valid:
                                value = round(totals.value_sum / usable, 2)
                                value_status = "estimated" if totals.reconstructed else "observed"
                            else:
                                value = None
                                value_status = "missing"
                                flags.add("insufficient_coverage")
                            fallback = station_first_cell.get(station_id)
                            source_id = totals.source_record_id or (
                                fallback.source_record_id
                                if fallback
                                else f"Standortdaten!{station_id}"
                            )
                            source_sheets = sorted(totals.source_sheets)
                            if not source_sheets and fallback:
                                source_sheets = [fallback.sheet]
                            records.append(
                                MobilityObservation(
                                    source_id=config.dataset_id,
                                    snapshot_id=raw.snapshot_id,
                                    source_record_id=source_id,
                                    record_id=f"window:{window.window_id}:{station_id}:{year}",
                                    geometry=station["geometry"],
                                    location_id=station_id,
                                    location_name=station["description"],
                                    mode="bicycle",
                                    metric_code="bicycle_mean_hourly_count",
                                    unit="bicycles_per_hour",
                                    aggregation="window",
                                    interval_start=f"{year:04d}-01-01",
                                    interval_end=f"{year + 1:04d}-01-01",
                                    timezone=timezone_name,
                                    measurement_status="measured",
                                    value=value,
                                    value_status=value_status,
                                    direction=station["direction"],
                                    window_id=window.window_id,
                                    expected_intervals=expected,
                                    observed_intervals=totals.observed,
                                    reconstructed_intervals=totals.reconstructed,
                                    implausible_intervals=totals.implausible,
                                    missing_intervals=missing,
                                    quality_flags=tuple(sorted(flags)),
                                    attributes={
                                        "install_date": station["install_date"].isoformat(),
                                        "covered_from": year_first.isoformat(),
                                        "covered_to": year_last.isoformat(),
                                        "holiday_adjustment": "none",
                                        "label": window.label,
                                        "source_sheets": ", ".join(source_sheets),
                                    },
                                )
                            )
            status_counts = defaultdict(int)
            for record in records:
                status_counts[record.value_status.value] += 1
            if unlegended_cell_count:
                examples = ", ".join(sorted(unlegended_examples)[:2])
                warnings.append(
                    f"{unlegended_cell_count} cells carry an undocumented source marking "
                    f"outside the spring DST hour (e.g. {examples}); their values "
                    f"({unlegended_numeric_count} numeric) were not used."
                )
            warnings.append(
                "value_status_counts: "
                + ", ".join(
                    f"{status}={status_counts[status]}"
                    for status in ("observed", "estimated", "missing")
                )
            )
            self._expected_station_count = sum(
                first_date is not None and station["install_date"] <= last_date
                for station in stations.values()
            )
            self._stations_with_records = {
                record.location_id for record in records if record.aggregation == "day"
            }
            self._duplicate_period_cells = duplicate_cell_count
            self._duplicate_conflict_cells = conflict_cell_count
            return NormalizedBatch(
                records=tuple(records),
                input_rows=input_rows,
                rejected_rows=len(rejected_rows),
                rejections=tuple(rejections),
                warnings=tuple(warnings),
            )
        finally:
            book.close()

    def validate(self, batch: NormalizedBatch) -> QualityReport:
        errors = list(batch.rejections)
        if not batch.records:
            errors.append("Bicycle counts workbook contains no observations")
        seen = set()
        stations = set()
        for record in batch.records:
            if not isinstance(record, MobilityObservation):
                errors.append("Bicycle counts batch contains a non-mobility record")
                continue
            if record.record_id in seen:
                errors.append(f"Duplicate mobility observation key: {record.record_id}")
            seen.add(record.record_id)
            stations.add(record.location_id)
        expected_stations = getattr(self, "_expected_station_count", None)
        if expected_stations is not None and len(stations) != expected_stations:
            errors.append(
                f"Expected records for {expected_stations} stations, found {len(stations)}"
            )
        return QualityReport(
            valid=not errors,
            input_rows=batch.input_rows,
            output_rows=len(batch.records),
            rejected_rows=batch.rejected_rows,
            warnings=batch.warnings,
            errors=tuple(errors),
        )

    @staticmethod
    def _years(value: object) -> set[int]:
        if not isinstance(value, list) or not value or any(type(year) is not int for year in value):
            raise ConfigError("years must be a nonempty list of integer years")
        return set(value)

    @staticmethod
    def _windows(value: object) -> tuple[_Window, ...]:
        if not isinstance(value, list) or not value:
            raise ConfigError("windows must be a nonempty list")
        result = []
        seen = set()
        for item in value:
            if not isinstance(item, dict):
                raise ConfigError("Each window must be an object")
            window_id, label, days = item.get("window_id"), item.get("label"), item.get("days")
            start, end = item.get("start_hour"), item.get("end_hour")
            if (
                not isinstance(window_id, str)
                or not window_id
                or not isinstance(label, str)
                or not label
                or days not in {"mon_fri", "sat", "sun", "all"}
                or type(start) is not int
                or type(end) is not int
                or not 0 <= start < end <= 24
                or window_id in seen
            ):
                raise ConfigError("Invalid or duplicate bicycle-count window")
            result.append(_Window(window_id, label, days, start, end))
            seen.add(window_id)
        return tuple(result)

    @staticmethod
    def _check_legend(sheet, options: dict) -> None:
        legend_fills = options.get("legend_fills")
        if not isinstance(legend_fills, dict):
            raise ConfigError("legend_fills must review the source legend")
        for coordinate in ("B2", "B3", "B4"):
            cell = sheet[coordinate]
            if not _matches_color(_color_signature(cell), legend_fills.get(coordinate, {})):
                raise SchemaError(f"Legend fill changed at Legende!{coordinate}")

    @staticmethod
    def _read_stations(sheet) -> dict[str, dict]:
        expected = (
            "Zählstelle",
            "Beschreibung - Fahrtrichtung",
            "Breitengrad",
            "Längengrad",
            "Installationsdatum",
        )
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if headers[: len(expected)] != expected or any(
            not _empty(value) for value in headers[len(expected) :]
        ):
            raise SchemaError("Standortdaten header changed")
        stations = {}
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if all(_empty(value) for value in row):
                continue
            station_id, description, latitude, longitude, installation = row[:5]
            if not isinstance(station_id, str) or not station_id.strip():
                raise SchemaError(f"Invalid station ID at Standortdaten!A{row_number}")
            if station_id in stations:
                raise SchemaError(f"Duplicate station ID in Standortdaten: {station_id}")
            if (
                not isinstance(description, str)
                or isinstance(latitude, bool)
                or not isinstance(latitude, (int, float))
                or isinstance(longitude, bool)
                or not isinstance(longitude, (int, float))
                or not isinstance(installation, (datetime, date))
            ):
                raise SchemaError(f"Invalid station metadata at Standortdaten row {row_number}")
            install_date = (
                installation.date() if isinstance(installation, datetime) else installation
            )
            tail = _normalized_space(description).rsplit(" ", 1)[-1]
            direction = _DIRECTIONS.get(tail)
            base_name = (
                _normalized_space(description).rsplit(" ", 1)[0]
                if direction is not None
                else _normalized_space(description)
            )
            stations[station_id] = {
                "description": description,
                "latitude": float(latitude),
                "longitude": float(longitude),
                "install_date": install_date,
                "direction": direction,
                "base_name": base_name,
                "geometry": {
                    "type": "Point",
                    "coordinates": [float(longitude), float(latitude)],
                },
            }
        if not stations:
            raise SchemaError("Standortdaten contains no stations")
        return stations

    @staticmethod
    def _station_flags(stations: dict[str, dict]) -> tuple[dict[str, set[str]], list]:
        from pyproj import Transformer

        transformer = Transformer.from_crs("EPSG:4326", "EPSG:25833", always_xy=True)
        flags = {station_id: set() for station_id in stations}
        projected = {
            station_id: transformer.transform(station["longitude"], station["latitude"])
            for station_id, station in stations.items()
        }
        pairs = []
        station_ids = sorted(stations)
        for index, first in enumerate(station_ids):
            for second in station_ids[index + 1 :]:
                if stations[first]["base_name"] == stations[second]["base_name"]:
                    continue
                distance = math.dist(projected[first], projected[second])
                if distance < 25:
                    flags[first].add("coordinates_unverified")
                    flags[second].add("coordinates_unverified")
                    pairs.append((first, second, distance))
        return flags, pairs

    @staticmethod
    def _resolve_alias(station_id: str, year: int, aliases: dict, stations: dict) -> str:
        alias = aliases.get(station_id)
        canonical = station_id
        if alias is not None:
            if not isinstance(alias, dict) or not isinstance(alias.get("canonical"), str):
                raise ConfigError(f"Invalid station alias configuration for {station_id}")
            active_years = alias.get("years", "all")
            if active_years == "all" or (isinstance(active_years, list) and year in active_years):
                canonical = alias["canonical"]
        if canonical not in stations:
            raise SchemaError(f"Unknown station header after aliasing: {station_id}")
        return canonical

    def _read_year_headers(
        self,
        sheet,
        sheet_name: str,
        sheet_year: int,
        stations: dict[str, dict],
        aliases: dict,
    ) -> dict[str, int]:
        max_column = sheet.max_column
        cells = next(sheet.iter_rows(min_row=1, max_row=1, max_col=max_column))
        if _normalized_space(cells[0].value) != "Zählstelle Inbetriebnahme":
            raise SchemaError(f"Year sheet header changed at {sheet_name}!A1")
        headers = {}
        canonical_ids = set()
        for column_number, cell in enumerate(cells[1:], 2):
            if cell.value is None:
                continue
            match = _HEADER.fullmatch(_normalized_space(cell.value))
            if not match:
                raise SchemaError(
                    f"Invalid counter header at {sheet_name}!{get_column_letter(column_number)}1"
                )
            source_id, install_text = match.groups()
            canonical = self._resolve_alias(source_id, sheet_year, aliases, stations)
            if canonical in canonical_ids:
                raise SchemaError(f"Station appears twice in {sheet_name}: {canonical}")
            canonical_ids.add(canonical)
            try:
                header_install_date = datetime.strptime(install_text, "%d.%m.%Y").date()
            except ValueError as exc:
                raise SchemaError(
                    f"Invalid installation date in {sheet_name}!{get_column_letter(column_number)}1"
                ) from exc
            if header_install_date != stations[canonical]["install_date"]:
                stations[canonical].setdefault("install_date_mismatch", True)
            direction_suffix = source_id.rsplit("-", 1)[-1]
            id_direction = _ID_DIRECTIONS.get(direction_suffix) if "-" in source_id else None
            if id_direction != stations[canonical]["direction"] and id_direction is not None:
                stations[canonical].setdefault("direction_label_conflict", True)
            headers[canonical] = column_number
        if not headers:
            raise SchemaError(f"Year sheet contains no counter columns: {sheet_name}")
        return headers

    @staticmethod
    def _cell_fill(cell, options: dict) -> str:
        signature = _color_signature(cell)
        if signature is None:
            return "none"
        fills = {
            "flagged_implausible": options["legend_fills"]["B2"],
            "missing": options["legend_fills"]["B3"],
            "reconstructed": options["legend_fills"]["B4"],
            "dst_gap": options.get("dst_gap_fill", {"type": "theme", "value": 7, "tint": 0.6}),
        }
        for fill_class, expected in fills.items():
            if _matches_color(signature, expected):
                return fill_class
        raise SchemaError(f"Unknown workbook fill colour: {signature}")

    @staticmethod
    def _count_value(value: object) -> int | None:
        if value is None or (isinstance(value, str) and not value.strip()):
            return None
        if isinstance(value, bool) or type(value) is not int:
            raise ValueError(f"count must be a nonnegative integer, found {value!r}")
        if value < 0:
            raise ValueError(f"count must be a nonnegative integer, found {value!r}")
        return value

    @staticmethod
    def _value_class(fill_class: str, value: int | None) -> str:
        if fill_class == "flagged_implausible":
            return "implausible"
        if fill_class == "missing":
            return "missing"
        if fill_class == "reconstructed":
            return "reconstructed" if value is not None else "missing"
        if fill_class == "dst_gap":
            return "dst_gap"
        if fill_class == "unlegended_source_marking":
            return "unlegended_source_marking"
        return "observed" if value is not None else "missing"

    @staticmethod
    def _spring_gap_day(day: date, zone: ZoneInfo) -> bool:
        midnight = datetime.combine(day, time.min, tzinfo=zone)
        next_midnight = datetime.combine(day + timedelta(days=1), time.min, tzinfo=zone)
        return next_midnight.utcoffset() > midnight.utcoffset()

    @staticmethod
    def _autumn_change_day(day: date, zone: ZoneInfo) -> bool:
        midnight = datetime.combine(day, time.min, tzinfo=zone)
        next_midnight = datetime.combine(day + timedelta(days=1), time.min, tzinfo=zone)
        return next_midnight.utcoffset() < midnight.utcoffset()

    @staticmethod
    def _window_day_matches(window: _Window, day: date) -> bool:
        return (
            window.days == "all"
            or window.days == "mon_fri"
            and day.weekday() < 5
            or window.days == "sat"
            and day.weekday() == 5
            or window.days == "sun"
            and day.weekday() == 6
        )

    @classmethod
    def _window_expected(cls, first: date, last: date, window: _Window, zone: ZoneInfo) -> int:
        expected = 0
        for day in cls._date_range(first, last):
            if not cls._window_day_matches(window, day):
                continue
            expected += window.end_hour - window.start_hour
            if cls._spring_gap_day(day, zone) and window.start_hour <= 2 < window.end_hour:
                expected -= 1
        return expected

    @classmethod
    def _apply_cell(
        cls,
        cell: _SourceCell,
        station_id: str,
        stamp: datetime,
        stations: dict[str, dict],
        windows: tuple[_Window, ...],
        zone: ZoneInfo,
        day_totals: dict,
        window_totals: dict,
        special_flags: dict,
        conflict: bool = False,
    ) -> None:
        day = stamp.date()
        station = stations[station_id]
        if day < station["install_date"]:
            return
        is_spring_hour = cls._spring_gap_day(day, zone) and stamp.hour == 2
        if cell.value_class == "dst_gap":
            if not is_spring_hour:
                raise SchemaError(
                    f"Unexpected DST-gap fill at {cell.sheet}!{cell.source_record_id}"
                )
            return
        if is_spring_hour:
            if cell.value_class == "observed":
                special_flags[(station_id, day)].add("dst_spring_gap_value_dropped")
            return
        totals = day_totals[(station_id, day)]
        totals.add(cell, conflict)
        for window in windows:
            if window.start_hour <= stamp.hour < window.end_hour and cls._window_day_matches(
                window, day
            ):
                window_totals[(station_id, day.year, window.window_id)].add(cell, conflict)

    @staticmethod
    def _date_range(first: date, last: date):
        day = first
        while day <= last:
            yield day
            day += timedelta(days=1)
