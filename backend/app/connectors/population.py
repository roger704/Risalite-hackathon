"""Reviewed tabular population imports; no guessed schemas or overlapping cohort sums."""

from __future__ import annotations

import csv
import math
import re
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from app.connectors.base import FileConnector
from app.domain import (
    AreaObservation,
    ConfigError,
    DatasetConfig,
    NormalizedBatch,
    QualityReport,
    RawBundle,
    SchemaError,
    ValidationError,
)


def _label(value: object) -> str:
    return re.sub(r"\s+", "", str(value or ""))


class PopulationConnector(FileConnector):
    adapter_version = "1.0.0"

    def normalize(self, raw: RawBundle, config: DatasetConfig) -> NormalizedBatch:
        if len(raw.assets) != 1:
            raise SchemaError("Population import requires exactly one reviewed tabular asset")
        options = config.parser_options
        if not config.reference_period:
            raise ConfigError("Population reference_period is required")
        metrics = options.get("metrics", [])
        self._check_metrics(metrics)
        path = Path(raw.assets[0].path)
        with path.open("rb") as handle:
            prefix = handle.read(512).lstrip().lower()
        if prefix.startswith((b"<!doctype html", b"<html")) or b"<html" in prefix:
            raise SchemaError("Population resource returned HTML instead of a table")
        fmt = options.get("format")
        if fmt == "xlsx":
            rows = self._xlsx_rows(path, options)
        elif fmt == "csv":
            rows = self._csv_rows(path, options)
        else:
            raise ConfigError("Population format must explicitly be xlsx or csv")
        records = []
        rejected = []
        input_rows = 0
        for row_number, row in rows:
            input_rows += 1
            try:
                code = self._area_code(row, options)
                observations = []
                for metric in metrics:
                    column = metric["column"]
                    value, status, flags = self._value(row[column], options)
                    lower, upper = metric.get("age_min"), metric.get("age_max_exclusive")
                    key = (
                        f"{code}:{config.reference_period}:{metric['metric_code']}:{lower}:{upper}"
                    )
                    observations.append(
                        AreaObservation(
                            source_id=config.dataset_id,
                            snapshot_id=raw.snapshot_id,
                            source_record_id=f"{options.get('worksheet', 'csv')}!{column}{row_number}",
                            record_id=key,
                            area_code=code,
                            hierarchy=options.get("hierarchy", "planungsraum"),
                            boundary_edition=config.boundary_edition,
                            metric_code=metric["metric_code"],
                            value=value,
                            unit=metric["unit"],
                            reference_period=config.reference_period,
                            value_status=status,
                            age_min=lower,
                            age_max_exclusive=upper,
                            quality_flags=flags,
                        )
                    )
                if options.get("check_age_total"):
                    total = next(o for o in observations if o.metric_code == "population_total")
                    bands = [o for o in observations if o.metric_code == "population_age_band"]
                    if (
                        total.value is not None
                        and bands
                        and all(o.value is not None for o in bands)
                    ):
                        if sum(o.value for o in bands) != total.value:
                            raise ValueError(
                                "nonoverlapping age bands do not sum to population total"
                            )
                records.extend(observations)
            except (ValueError, KeyError, ValidationError) as exc:
                rejected.append(f"Row {row_number}: {exc}")
        expected = options.get("expected_area_count")
        if expected is not None and input_rows != expected:
            rejected.append(f"Expected {expected} area rows, found {input_rows}")
        return NormalizedBatch(
            records=tuple(records),
            input_rows=input_rows,
            rejected_rows=input_rows - len(records) // len(metrics),
            rejections=tuple(rejected),
            warnings=tuple(options.get("limitations", ())),
        )

    @staticmethod
    def _check_metrics(metrics: list[dict]) -> None:
        if not metrics:
            raise ConfigError("Population metrics require an explicit reviewed mapping")
        keys = set()
        bands: dict[tuple[str, str], list[tuple[int, float]]] = {}
        for metric in metrics:
            if not all(metric.get(key) for key in ("column", "metric_code", "unit")):
                raise ConfigError("Every population metric needs column, metric_code, and unit")
            lo, hi = metric.get("age_min"), metric.get("age_max_exclusive")
            key = (metric["metric_code"], lo, hi)
            if key in keys:
                raise ConfigError("Duplicate population metric mapping")
            keys.add(key)
            if (
                lo is None
                and hi is not None
                or lo is not None
                and (
                    type(lo) is not int
                    or lo < 0
                    or hi is not None
                    and (type(hi) is not int or hi <= lo)
                )
            ):
                raise ConfigError("Invalid age band bounds")
            if lo is not None:
                group = bands.setdefault((metric["metric_code"], metric["unit"]), [])
                upper = hi if hi is not None else math.inf
                if any(lo < old_hi and old_lo < upper for old_lo, old_hi in group):
                    raise ConfigError(
                        "Overlapping population age bands are not permitted within a metric"
                    )
                group.append((lo, upper))

    @staticmethod
    def _csv_rows(path: Path, options: dict) -> list[tuple[int, dict]]:
        encoding = options.get("encoding")
        delimiter = options.get("delimiter")
        if not encoding or not isinstance(delimiter, str) or len(delimiter) != 1:
            raise ConfigError("CSV requires reviewed encoding and one-character delimiter")
        try:
            with path.open(encoding=encoding, newline="") as handle:
                for _ in range(options.get("skip_rows", 0)):
                    next(handle)
                reader = csv.DictReader(handle, delimiter=delimiter)
                fields = reader.fieldnames or []
                required = set(options.get("required_fields", []))
                required.update(m["column"] for m in options["metrics"])
                required.update(options.get("area_code_columns", []))
                if len(fields) != len(set(fields)) or not required.issubset(fields):
                    raise SchemaError(f"Population CSV fields changed; required {sorted(required)}")
                result = []
                for row in reader:
                    if None in row or any(value is None for value in row.values()):
                        raise SchemaError("Population CSV row width does not match reviewed header")
                    result.append((reader.line_num + options.get("skip_rows", 0), row))
                return result
        except (UnicodeError, csv.Error, StopIteration) as exc:
            raise SchemaError(
                f"Population CSV cannot be parsed with reviewed settings: {exc}"
            ) from exc

    @staticmethod
    def _xlsx_rows(path: Path, options: dict) -> list[tuple[int, dict]]:
        from openpyxl import load_workbook
        from openpyxl.utils import column_index_from_string

        try:
            with ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > options.get(
                    "max_unpacked_bytes", 30_000_000
                ):
                    raise SchemaError("Population workbook exceeds unpacked size budget")
            book = load_workbook(path, read_only=True, data_only=False)
        except (BadZipFile, OSError, ValueError) as exc:
            raise SchemaError("Population resource is not a readable XLSX workbook") from exc
        try:
            worksheet = options.get("worksheet")
            if worksheet not in book.sheetnames:
                raise SchemaError(f"Population worksheet {worksheet!r} is missing")
            sheet = book[worksheet]
            checks = options.get("expected_cells", {})
            if not checks:
                raise ConfigError("XLSX requires reviewed title and column-header cell checks")
            for coordinate, expected in checks.items():
                if _label(sheet[coordinate].value) != _label(expected):
                    raise SchemaError(f"Population XLSX schema changed at {worksheet}!{coordinate}")
            columns = set(options["area_code_columns"])
            columns.update(m["column"] for m in options["metrics"])
            indexes = {column: column_index_from_string(column) - 1 for column in columns}
            first = options.get("first_data_row", 1)
            result = []
            for number, values in enumerate(
                sheet.iter_rows(min_row=first, max_col=max(indexes.values()) + 1, values_only=True),
                first,
            ):
                row = {column: values[index] for column, index in indexes.items()}
                # Reviewed report layout contains district headers, blank rows, and
                # aggregate totals with no district code. They are not PLR observations.
                if row[options["area_code_columns"][0]] is None:
                    continue
                result.append((number, row))
            return result
        finally:
            book.close()

    @staticmethod
    def _area_code(row: dict, options: dict) -> str:
        components = []
        widths = options.get("area_code_widths", [8])
        columns = options.get("area_code_columns", [])
        if len(columns) != len(widths) or not columns:
            raise ConfigError("Population area code columns and widths must match")
        for column, width in zip(columns, widths):
            value = row[column]
            if type(value) is int and options.get("allow_numeric_code_components", False):
                value = str(value).zfill(width)
            if not isinstance(value, str) or not re.fullmatch(rf"[0-9]{{{width}}}", value.strip()):
                raise ValueError(f"invalid area code component {column}: {value!r}")
            components.append(value.strip())
        return "".join(components)

    @staticmethod
    def _value(value: object, options: dict) -> tuple[int | float | None, str, tuple[str, ...]]:
        text = "" if value is None else str(value).strip()
        sentinels = options.get("sentinels", {})
        if text in sentinels:
            status = sentinels[text]
            if status == "zero":
                return 0, "observed", (f"source_marker:{text}",)
            if status not in {"suppressed", "missing", "not_applicable"}:
                raise ConfigError(f"Unsupported sentinel status {status}")
            return None, status, (f"source_marker:{text}",)
        thousands, decimal = (
            options.get("thousands_separator", ""),
            options.get("decimal_separator", "."),
        )
        if thousands:
            text = text.replace(thousands, "")
        if decimal != ".":
            text = text.replace(decimal, ".")
        if not re.fullmatch(r"[0-9]+(?:\.[0-9]+)?", text):
            raise ValueError(f"unknown population value {value!r}; review source legend")
        number = float(text)
        if not math.isfinite(number):
            raise ValueError("Population count must be finite")
        if options.get("integer_counts", True) and not number.is_integer():
            raise ValueError("Population counts must be integers")
        return int(number) if number.is_integer() else number, "observed", ()

    def validate(self, batch: NormalizedBatch) -> QualityReport:
        errors = list(batch.rejections)
        seen = set()
        if not batch.records:
            errors.append("Population import contains no usable observations")
        for record in batch.records:
            if not isinstance(record, AreaObservation):
                errors.append("Population batch contains a non-observation record")
                continue
            key = (
                record.area_code,
                record.hierarchy,
                record.boundary_edition,
                record.reference_period,
                record.metric_code,
                record.age_min,
                record.age_max_exclusive,
            )
            if key in seen:
                errors.append(f"Duplicate population observation key: {key}")
            seen.add(key)
        return QualityReport(
            valid=not errors,
            input_rows=batch.input_rows,
            output_rows=len(batch.records),
            rejected_rows=batch.rejected_rows,
            warnings=batch.warnings,
            errors=tuple(errors),
        )
