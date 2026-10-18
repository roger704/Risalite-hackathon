"""Pinned Zensus 2022 INSPIRE population-grid import."""

from __future__ import annotations

import csv
import math
import re
from io import BytesIO, TextIOWrapper
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile, ZipFile, is_zipfile

from pyproj import Transformer
from shapely.geometry import box, shape
from shapely.ops import transform, unary_union
from shapely.prepared import prep

from app.connectors.base import FileConnector
from app.domain import (
    Area,
    ConfigError,
    DatasetConfig,
    GridObservation,
    IncompatibleSnapshotError,
    NormalizedBatch,
    QualityReport,
    RawBundle,
    SchemaError,
    ValidationError,
)
from app.geo.inspire_grid import cell_geojson_wgs84, cell_polygon, parse_grid_id


class ZensusGridConnector(FileConnector):
    adapter_version = "1.0.0"

    def normalize(self, raw: RawBundle, config: DatasetConfig) -> NormalizedBatch:
        if len(raw.assets) != 1:
            raise SchemaError("Zensus grid import requires exactly one ZIP asset")
        if not config.reference_period:
            raise ConfigError("Zensus grid reference_period is required")
        asset = raw.assets[0]
        options = config.parser_options
        expected_sha256 = options.get("expected_sha256")
        if expected_sha256 and asset.sha256.lower() != expected_sha256.lower():
            raise SchemaError(
                "Zensus source file changed; review the new edition before updating the pin"
            )
        path = Path(asset.path)
        if not is_zipfile(path):
            raise SchemaError("Zensus source is not a ZIP archive")
        source_bbox = _bbox(options.get("source_bbox"))
        resolution_m = options.get("resolution_m")
        minimum_cells = options.get("min_expected_cells", 0)
        encoding = options.get("encoding")
        delimiter = options.get("delimiter")
        if type(resolution_m) is not int or resolution_m <= 0:
            raise SchemaError("Zensus resolution_m must be a positive integer")
        if type(minimum_cells) is not int or minimum_cells < 0:
            raise SchemaError("Zensus min_expected_cells must be a nonnegative integer")
        if not isinstance(encoding, str) or not encoding:
            raise SchemaError("Zensus CSV encoding must be explicit")
        if not isinstance(delimiter, str) or len(delimiter) != 1:
            raise SchemaError("Zensus CSV delimiter must be one character")
        expected_header = options.get("expected_header")
        if not isinstance(expected_header, list) or any(
            not isinstance(value, str) for value in expected_header
        ):
            raise SchemaError("Zensus expected_header must be a list of strings")
        archive_member = options.get("archive_member")
        description_member = options.get("description_member")
        if not isinstance(archive_member, str) or not isinstance(description_member, str):
            raise SchemaError("Zensus archive_member and description_member are required")
        max_unpacked_bytes = options.get("max_unpacked_bytes")
        if type(max_unpacked_bytes) is not int or max_unpacked_bytes <= 0:
            raise SchemaError("Zensus max_unpacked_bytes must be a positive integer")

        records: list[GridObservation] = []
        rejection_messages: list[str] = []
        rejected_rows = 0
        national_rows = 0

        def reject(line_number: int, reason: str) -> None:
            nonlocal rejected_rows
            rejected_rows += 1
            if len(rejection_messages) < 20:
                rejection_messages.append(f"line {line_number}: {reason}")

        try:
            with ZipFile(path) as archive:
                members = archive.infolist()
                names = [member.filename for member in members]
                for name in names:
                    _validate_member_path(name)
                if names.count(archive_member) != 1 or names.count(description_member) != 1:
                    raise SchemaError("Zensus ZIP is missing or duplicates a required member")
                for member in members:
                    if member.file_size > max_unpacked_bytes:
                        raise SchemaError(
                            f"Zensus ZIP member exceeds max_unpacked_bytes: {member.filename}"
                        )
                description = archive.read(description_member)
                _check_description(description, options.get("expected_texts", []))
                info = archive.getinfo(archive_member)
                with (
                    archive.open(info) as binary,
                    TextIOWrapper(binary, encoding=encoding, newline="") as source,
                ):
                    reader = csv.reader(source, delimiter=delimiter)
                    header = next(reader, None)
                    if header != expected_header:
                        raise SchemaError("Zensus CSV header differs from expected_header")
                    for row in reader:
                        national_rows += 1
                        line_number = reader.line_num
                        if len(row) != len(expected_header):
                            reject(line_number, "wrong column count")
                            continue
                        try:
                            cell = parse_grid_id(row[0].strip())
                            if cell.resolution_m != resolution_m:
                                raise ValueError("wrong grid resolution")
                            east_mp = _integer(row[1], "x_mp")
                            north_mp = _integer(row[2], "y_mp")
                            if (
                                east_mp != cell.east + cell.resolution_m / 2
                                or north_mp != cell.north + cell.resolution_m / 2
                            ):
                                raise ValueError("cell centre does not match grid ID")
                            if not _intersects_bbox(
                                cell.east, cell.north, cell.resolution_m, source_bbox
                            ):
                                continue
                            value, status, flags = _value(row[3], options.get("sentinels", {}))
                            records.append(
                                GridObservation(
                                    source_id=config.dataset_id,
                                    snapshot_id=raw.snapshot_id,
                                    source_record_id=cell.grid_id,
                                    record_id=(
                                        f"{cell.grid_id}:{options['metric_code']}:"
                                        f"{config.reference_period}"
                                    ),
                                    geometry=cell_geojson_wgs84(cell),
                                    grid_id=cell.grid_id,
                                    grid_crs=options["grid_crs"],
                                    resolution_m=cell.resolution_m,
                                    metric_code=options["metric_code"],
                                    value=value,
                                    unit=options["unit"],
                                    reference_period=config.reference_period,
                                    value_status=status,
                                    quality_flags=flags,
                                )
                            )
                        except (KeyError, TypeError, ValueError, ValidationError) as exc:
                            reject(line_number, str(exc))
        except SchemaError:
            raise
        except (BadZipFile, OSError, UnicodeError, csv.Error, KeyError, ValueError) as exc:
            raise SchemaError(f"Unable to read Zensus ZIP or CSV: {exc}") from exc

        if len(records) < minimum_cells:
            rejection_messages.append(
                f"Expected at least {minimum_cells} cells, found {len(records)}"
            )
        if rejected_rows > 20:
            rejection_messages.append(f"{rejected_rows - 20} additional rejected rows omitted")
        published_sum = sum(record.value for record in records if record.value is not None)
        warnings = tuple(options.get("limitations", ())) + (
            f"Scanned {national_rows} rows of {archive_member}; kept {len(records)} cells "
            f"intersecting source_bbox {list(source_bbox)} (EPSG:3035)",
            "Sum of published cell values in the extract: "
            f"{published_sum} persons (cells outside Berlin included)",
        )
        return NormalizedBatch(
            records=tuple(records),
            input_rows=len(records) + rejected_rows,
            rejected_rows=rejected_rows,
            rejections=tuple(rejection_messages),
            warnings=warnings,
        )

    def validate(self, batch: NormalizedBatch) -> QualityReport:
        errors = list(batch.rejections)
        if not batch.records:
            errors.append("Zensus grid snapshot contains no cells")
        if batch.rejected_rows:
            errors.append("Zensus grid snapshot contains rejected rows")
        seen: set[str] = set()
        resolutions = set()
        grid_crs_values = set()
        periods = set()
        metrics = set()
        for record in batch.records:
            if not isinstance(record, GridObservation):
                errors.append("Zensus grid batch contains a non-grid observation")
                continue
            if record.grid_id in seen:
                errors.append(f"Duplicate Zensus grid ID: {record.grid_id}")
            seen.add(record.grid_id)
            resolutions.add(record.resolution_m)
            grid_crs_values.add(record.grid_crs)
            periods.add(record.reference_period)
            metrics.add(record.metric_code)
        for label, values in (
            ("resolution", resolutions),
            ("grid CRS", grid_crs_values),
            ("reference period", periods),
            ("metric code", metrics),
        ):
            if len(values) > 1:
                errors.append(f"Zensus grid snapshot mixes {label} values")
        return QualityReport(
            valid=not errors,
            input_rows=batch.input_rows,
            output_rows=len(batch.records),
            rejected_rows=batch.rejected_rows,
            warnings=batch.warnings,
            errors=tuple(errors),
        )

    def check_parent(self, config: DatasetConfig, records: tuple, parent_records: tuple) -> dict:
        areas = [record for record in parent_records if isinstance(record, Area)]
        if not areas:
            raise IncompatibleSnapshotError("Zensus grid requires Berlin LOR Area records")
        options = config.parser_options
        source_bbox = _bbox(options.get("source_bbox"))
        edge_buffer_m = options.get("edge_buffer_m", 0)
        if (
            isinstance(edge_buffer_m, bool)
            or not isinstance(edge_buffer_m, (int, float))
            or not math.isfinite(edge_buffer_m)
            or edge_buffer_m < 0
        ):
            raise SchemaError("Zensus edge_buffer_m must be a finite nonnegative number")
        to_25833 = Transformer.from_crs("EPSG:4326", "EPSG:25833", always_xy=True)
        to_3035 = Transformer.from_crs("EPSG:4326", "EPSG:3035", always_xy=True)
        berlin_wgs84 = unary_union([shape(area.geometry) for area in areas])
        berlin_25833 = transform(to_25833.transform, berlin_wgs84)
        berlin_3035 = transform(to_3035.transform, berlin_wgs84)
        buffered_boundary = berlin_3035.buffer(edge_buffer_m)
        if not box(*source_bbox).covers(buffered_boundary):
            raise IncompatibleSnapshotError(
                "Zensus source_bbox truncates Berlin or its configured edge buffer"
            )
        prepared_berlin = prep(berlin_25833)
        counts = {
            "cells_in_extract": len(records),
            "cells_intersecting_berlin": 0,
            "cells_within_berlin": 0,
            "cells_partially_in_berlin": 0,
            "cells_outside_berlin": 0,
        }
        intersecting_population = 0
        within_population = 0
        centroid_population = 0
        area_weighted_population = 0.0
        for record in records:
            if not isinstance(record, GridObservation):
                raise IncompatibleSnapshotError("Zensus grid contains a non-grid observation")
            cell = parse_grid_id(record.grid_id)
            polygon = cell_polygon(cell, "EPSG:25833")
            if prepared_berlin.intersects(polygon):
                counts["cells_intersecting_berlin"] += 1
                intersection = berlin_25833.intersection(polygon).area
                within = berlin_25833.covers(polygon)
                if within:
                    counts["cells_within_berlin"] += 1
                else:
                    counts["cells_partially_in_berlin"] += 1
                if record.value is not None:
                    intersecting_population += record.value
                    area_weighted_population += record.value * intersection / polygon.area
                    if within:
                        within_population += record.value
                if berlin_25833.covers(polygon.centroid) and record.value is not None:
                    centroid_population += record.value
            else:
                counts["cells_outside_berlin"] += 1
        official_total = dict(options.get("official_total", {}))
        official_value = official_total.get("value")
        ratio = (
            round(centroid_population / official_value, 4)
            if isinstance(official_value, (int, float))
            and not isinstance(official_value, bool)
            and official_value != 0
            else None
        )
        first = records[0] if records else None
        return {
            "parent_dataset": "berlin_lor",
            "grid_crs": first.grid_crs if first else options.get("grid_crs"),
            "resolution_m": first.resolution_m if first else options.get("resolution_m"),
            "reference_period": first.reference_period if first else config.reference_period,
            "metric_code": first.metric_code if first else options.get("metric_code"),
            "berlin_comparison": {
                **counts,
                "population_cells_intersecting_berlin": intersecting_population,
                "population_cells_within_berlin": within_population,
                "population_centroid_in_berlin": centroid_population,
                "population_area_weighted_in_berlin": round(area_weighted_population, 1),
                "official_total": official_total,
                "centroid_to_official_ratio": ratio,
            },
            "method": (
                "cell polygons rebuilt from INSPIRE IDs and intersected with the union of "
                "berlin_lor polygons in EPSG:25833"
            ),
            "limitations": [
                "Area-weighted values assume uniform density within each cell and are estimates.",
                "Grid sums differ from the official Berlin total because of CKM perturbation, "
                "unpublished cells and addresses missing from grid data.",
            ],
        }


def _bbox(value) -> tuple[float, float, float, float]:
    if (
        not isinstance(value, (list, tuple))
        or len(value) != 4
        or any(
            isinstance(item, bool) or not isinstance(item, (int, float)) or not math.isfinite(item)
            for item in value
        )
        or value[0] >= value[2]
        or value[1] >= value[3]
    ):
        raise SchemaError("Zensus parser requires a valid source_bbox")
    return tuple(value)


def _validate_member_path(name: str) -> None:
    normalized = name.replace("\\", "/")
    path = PurePosixPath(normalized)
    if (
        "\x00" in name
        or path.is_absolute()
        or any(part == ".." for part in path.parts)
        or re.match(r"^[A-Za-z]:", normalized)
    ):
        raise SchemaError(f"Unsafe path in Zensus ZIP: {name!r}")


def _check_description(data: bytes, expected_texts) -> None:
    if not isinstance(expected_texts, list) or any(
        not isinstance(text, str) or not text for text in expected_texts
    ):
        raise SchemaError("Zensus expected_texts must be a list of nonempty strings")
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            text_cells = [
                value
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows(values_only=True)
                for value in row
                if isinstance(value, str)
            ]
        finally:
            workbook.close()
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise SchemaError("Zensus description workbook is not readable") from exc
    joined = "\n".join(text_cells)
    for expected in expected_texts:
        if expected not in joined:
            raise SchemaError(f"Zensus description is missing expected text: {expected}")


def _integer(value: str, name: str) -> int:
    text = value.strip()
    if not re.fullmatch(r"-?\d+", text):
        raise ValueError(f"{name} is not an integer")
    return int(text)


def _intersects_bbox(east: int, north: int, resolution: int, bounds) -> bool:
    xmin, ymin, xmax, ymax = bounds
    return east < xmax and east + resolution > xmin and north < ymax and north + resolution > ymin


def _value(value: str, sentinels: dict) -> tuple[int | None, str, tuple[str, ...]]:
    text = value.strip()
    if re.fullmatch(r"\d+", text):
        return int(text), "observed", ()
    if text in sentinels:
        status = sentinels[text]
        if status == "zero":
            return 0, "observed", ("zero_or_set_to_zero",)
        if status in {"missing", "suppressed", "not_applicable"}:
            return None, status, ()
        raise ConfigError(f"Unsupported Zensus sentinel status: {status}")
    raise ValueError(f"unknown value token {value!r}")
