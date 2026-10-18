"""Pinned Berlin population forecast workbooks at their published geographies."""

from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import replace
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from app.connectors.base import FileConnector
from app.domain import (
    Area,
    ConfigError,
    DatasetConfig,
    ForecastObservation,
    IncompatibleSnapshotError,
    NormalizedBatch,
    QualityReport,
    RawBundle,
    SchemaError,
    ValidationError,
)

_YEAR_LABELS = tuple(range(2024, 2041))
_SCENARIO_COLUMNS = {"middle": 2, "upper": 4, "lower": 6}
_SCENARIO_LABELS = {
    "middle": "Mittlere Variante",
    "upper": "Obere Variante",
    "lower": "Untere Variante",
}
_AGE_LABELS = tuple(f"{age:02d} - unter {age + 1:02d}" for age in range(99)) + ("99 und älter",)


def _normalize_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _record_id(
    edition: str,
    scenario: str,
    hierarchy: str,
    area_code: str,
    period: str,
    metric: str,
    age_min: int | None,
    age_max_exclusive: int | None,
    unit: str,
) -> str:
    return (
        f"{edition}:{scenario}:{hierarchy}:{area_code}:{period}:{metric}:"
        f"{age_min}:{age_max_exclusive}:{unit}"
    )


def expected_change(base: ForecastObservation, target: ForecastObservation) -> dict:
    comparable = (
        "area_code",
        "hierarchy",
        "boundary_edition",
        "scenario",
        "forecast_edition",
        "metric_code",
        "unit",
        "age_min",
        "age_max_exclusive",
    )
    if any(getattr(base, key) != getattr(target, key) for key in comparable):
        raise ValidationError(
            "Forecast change requires matching geography, scenario, metric, and unit"
        )
    if base.value_status != "observed":
        raise ValidationError("Forecast change base must be observed")
    if target.value_status != "forecast" or base.reference_period != target.base_period:
        raise ValidationError("Forecast target must use the observed base period")
    if date.fromisoformat(target.reference_period) <= date.fromisoformat(base.reference_period):
        raise ValidationError("Forecast target must be later than its observed base")
    absolute = target.value - base.value
    percent = None if base.value == 0 else absolute / base.value * 100
    names = {"land": "Berlin", "bezirk": "Bezirk", "prognoseraum": "Prognoseraum"}
    return {
        "area_code": base.area_code,
        "hierarchy": base.hierarchy,
        "scenario": base.scenario,
        "base_period": base.reference_period,
        "target_period": target.reference_period,
        "unit": base.unit,
        "base_value": base.value,
        "target_value": target.value,
        "absolute_change": absolute,
        "percent_change": percent,
        "resolution_note": (
            f"Forecast for {names[base.hierarchy]} {base.area_code} ({base.scenario} variant); "
            "not a measured planning-area or address-level value."
        ),
    }


class PopulationForecastConnector(FileConnector):
    adapter_version = "1.0"

    async def fetch(self, request) -> RawBundle:
        config = request.config
        pins = self._pins(config)
        primary = await super().fetch(request)
        overview_pin = next(pin for pin in pins if pin["role"] == "scenario_overview")
        overview_config = replace(config, resource_url=overview_pin["url"])
        overview = await self.fetcher.download(
            overview_config, Path(request.staging_dir) / "scenario_overview.xlsx"
        )
        assets = (primary.assets[0], overview)
        for pin, asset in zip(pins, assets):
            if asset.sha256 != pin["sha256"]:
                raise ValidationError("Pinned edition changed; review before refresh")
            if asset.size_bytes != pin["size_bytes"]:
                raise ValidationError("Pinned edition changed; review before refresh")
        fingerprint = hashlib.sha256(
            json.dumps(
                {
                    "urls": [pin["url"] for pin in pins],
                    "pins": [
                        {key: pin[key] for key in ("role", "sha256", "size_bytes")} for pin in pins
                    ],
                    "source_edition": config.source_edition,
                    "adapter_version": self.adapter_version,
                },
                sort_keys=True,
            ).encode()
        ).hexdigest()
        return RawBundle(
            assets=assets,
            retrieved_at=primary.retrieved_at,
            snapshot_id=request.snapshot_id,
            request_fingerprint=fingerprint,
        )

    @staticmethod
    def _pins(config: DatasetConfig) -> list[dict]:
        pins = config.parser_options.get("assets")
        if not isinstance(pins, list) or len(pins) != 2:
            raise ConfigError("Population forecast requires two pinned workbook assets")
        roles = {pin.get("role") for pin in pins if isinstance(pin, dict)}
        if roles != {"small_area", "scenario_overview"}:
            raise ConfigError(
                "Population forecast asset pins need small_area and scenario_overview"
            )
        if any(
            not isinstance(pin.get(key), expected)
            for pin in pins
            if isinstance(pin, dict)
            for key, expected in (
                ("url", str),
                ("sha256", str),
                ("size_bytes", int),
                ("role", str),
            )
        ):
            raise ConfigError("Population forecast asset pins are incomplete")
        if any(not isinstance(pin, dict) for pin in pins):
            raise ConfigError("Population forecast asset pins are invalid")
        small = next(pin for pin in pins if pin["role"] == "small_area")
        if small["url"] != config.resource_url:
            raise ConfigError("Small-area asset pin must match resource_url")
        return [small, next(pin for pin in pins if pin["role"] == "scenario_overview")]

    def normalize(self, raw: RawBundle, config: DatasetConfig) -> NormalizedBatch:
        pins = self._pins(config)
        if len(raw.assets) != 2:
            raise ConfigError(
                "Population forecast requires both pinned workbooks; local single-file imports are unsupported"
            )
        by_url = {asset.resource_url: asset for asset in raw.assets}
        if len(by_url) != len(raw.assets) or set(by_url) != {pin["url"] for pin in pins}:
            raise SchemaError(
                "Population forecast assets do not match the two reviewed workbook URLs"
            )
        for pin in pins:
            asset = by_url[pin["url"]]
            if asset.sha256 != pin["sha256"] or asset.size_bytes != pin["size_bytes"]:
                raise ValidationError("Pinned edition changed; review before refresh")

        options = config.parser_options
        small_pin = next(pin for pin in pins if pin["role"] == "small_area")
        overview_pin = next(pin for pin in pins if pin["role"] == "scenario_overview")
        small_path = Path(by_url[small_pin["url"]].path)
        overview_path = Path(by_url[overview_pin["url"]].path)
        self._check_workbook_size(small_path, options)
        self._check_workbook_size(overview_path, options)
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        records: list[ForecastObservation] = []
        rejections: list[str] = []
        input_rows = rejected_rows = 0
        total_values: dict[tuple[str, str, int], float] = {}
        district_codes: set[str] = set()
        pgr_codes: set[str] = set()
        max_age_discrepancy = 0.0
        age_checks = 0

        def emit(
            *,
            sheet: str,
            row_number: int,
            column: int,
            hierarchy: str,
            area_code: str,
            year: int,
            metric_code: str,
            value: object,
            age_min: int | None = None,
            age_max_exclusive: int | None = None,
            unit: str = "persons",
            scenario: str = "middle",
            flags: tuple[str, ...] = (),
        ) -> float | None:
            nonlocal input_rows, rejected_rows
            input_rows += 1
            if (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
                or value < 0
            ):
                rejected_rows += 1
                rejections.append(
                    f"{sheet}!{get_column_letter(column)}{row_number}: "
                    f"non-numeric or empty population value {value!r}"
                )
                return None
            period = f"{year}-12-31"
            record = ForecastObservation(
                source_id=config.dataset_id,
                snapshot_id=raw.snapshot_id,
                source_record_id=f"{sheet}!{get_column_letter(column)}{row_number}",
                record_id=_record_id(
                    options["forecast_edition"],
                    scenario,
                    hierarchy,
                    area_code,
                    period,
                    metric_code,
                    age_min,
                    age_max_exclusive,
                    unit,
                ),
                area_code=area_code,
                hierarchy=hierarchy,
                boundary_edition=config.boundary_edition,
                metric_code=metric_code,
                value=value,
                unit=unit,
                reference_period=period,
                value_status="observed" if period == options["base_period"] else "forecast",
                scenario=scenario,
                forecast_edition=options["forecast_edition"],
                base_period=options["base_period"],
                age_min=age_min,
                age_max_exclusive=age_max_exclusive,
                quality_flags=flags,
            )
            records.append(record)
            return float(value)

        try:
            book = load_workbook(small_path, read_only=True, data_only=True)
        except (BadZipFile, OSError, ValueError) as exc:
            raise SchemaError(
                "Small-area forecast resource is not a readable XLSX workbook"
            ) from exc
        try:
            expected_sheets = options["expected_sheet_names"]["small_area"]
            if book.sheetnames != expected_sheets:
                raise SchemaError(
                    f"Small-area workbook sheets changed: expected {expected_sheets}, "
                    f"found {book.sheetnames}"
                )
            if "00_Berlin" not in book.sheetnames:
                raise SchemaError("Small-area workbook is missing 00_Berlin")
            for sheet_name in book.sheetnames:
                if sheet_name == options.get("overview_sheet", "00_Übersicht"):
                    continue
                sheet = book[sheet_name]
                rows = list(
                    sheet.iter_rows(
                        min_row=1, max_row=109, max_col=sheet.max_column, values_only=True
                    )
                )
                self._check_small_area_header(sheet_name, rows, options, config)
                pgr_starts: list[tuple[int, str]] = []
                if sheet_name == "00_Berlin":
                    areas = [(3, "land", "11")]
                else:
                    district_code = sheet_name[:2]
                    district_codes.add(district_code)
                    areas = [(3, "bezirk", district_code)]
                    for column, label in enumerate(rows[6], start=1):
                        if label is None:
                            continue
                        if column < 21 or not isinstance(label, str):
                            raise SchemaError(f"Unexpected row-7 label in {sheet_name}!{column}")
                        match = re.fullmatch(r"(\d{4}) - .+", label)
                        if not match or match[1][:2] != district_code:
                            raise SchemaError(
                                f"Invalid Prognoseraum label in {sheet_name}!{column}7: {label!r}"
                            )
                        pgr_starts.append((column, match[1]))
                    if not pgr_starts:
                        raise SchemaError(f"No Prognoseraum blocks found in {sheet_name}")
                    areas.extend((column, "prognoseraum", code) for column, code in pgr_starts)
                    pgr_codes.update(code for _, code in pgr_starts)

                for column_start, hierarchy, area_code in areas:
                    area_ages: dict[int, list[float | None]] = {year: [] for year in _YEAR_LABELS}
                    for year_offset, year in enumerate(_YEAR_LABELS):
                        column = column_start + year_offset
                        total = emit(
                            sheet=sheet_name,
                            row_number=109,
                            column=column,
                            hierarchy=hierarchy,
                            area_code=area_code,
                            year=year,
                            metric_code="population_total",
                            value=rows[108][column - 1] if column <= len(rows[108]) else None,
                        )
                        key = (hierarchy, area_code, year)
                        if total is not None:
                            total_values[key] = total
                        for age in range(100):
                            source_row = 9 + age
                            value = (
                                rows[source_row - 1][column - 1]
                                if column <= len(rows[source_row - 1])
                                else None
                            )
                            parsed = emit(
                                sheet=sheet_name,
                                row_number=source_row,
                                column=column,
                                hierarchy=hierarchy,
                                area_code=area_code,
                                year=year,
                                metric_code="population_age",
                                value=value,
                                age_min=age,
                                age_max_exclusive=age + 1 if age < 99 else None,
                            )
                            area_ages[year].append(parsed)
                        age_sum = (
                            sum(area_ages[year])
                            if all(value is not None for value in area_ages[year])
                            else None
                        )
                        if total is not None and age_sum is not None:
                            age_checks += 1
                            discrepancy = abs(age_sum - total)
                            max_age_discrepancy = max(max_age_discrepancy, discrepancy)
                            if discrepancy > 1.0:
                                raise ValidationError(
                                    f"Age values for {area_code} in {year} differ from total "
                                    f"by {discrepancy:.6f} persons"
                                )
        finally:
            book.close()

        expected_counts = options["expected_area_counts"]
        if len(district_codes) != expected_counts["bezirk"]:
            raise SchemaError(
                f"Expected {expected_counts['bezirk']} districts, found {len(district_codes)}"
            )
        if len(pgr_codes) != expected_counts["prognoseraum"]:
            raise SchemaError(
                f"Expected {expected_counts['prognoseraum']} Prognoseräume, found {len(pgr_codes)}"
            )

        max_pgr_district = max_district_berlin = 0.0
        pgr_district_violations = district_berlin_violations = 0
        berlin_values = {year: total_values.get(("land", "11", year)) for year in _YEAR_LABELS}
        for year in _YEAR_LABELS:
            for district in district_codes:
                district_value = total_values.get(("bezirk", district, year))
                pgr_values = [
                    total_values.get(("prognoseraum", code, year))
                    for code in sorted(pgr_codes)
                    if code[:2] == district
                ]
                if district_value is not None and all(value is not None for value in pgr_values):
                    discrepancy = abs(sum(pgr_values) - district_value)
                    max_pgr_district = max(max_pgr_district, discrepancy)
                    pgr_district_violations += discrepancy > 1.0
            district_values = [
                total_values.get(("bezirk", district, year)) for district in sorted(district_codes)
            ]
            if berlin_values[year] is not None and all(
                value is not None for value in district_values
            ):
                discrepancy = abs(sum(district_values) - berlin_values[year])
                max_district_berlin = max(max_district_berlin, discrepancy)
                district_berlin_violations += discrepancy > 1.0

        self._normalize_overview(overview_path, raw, config, options, emit, rejections, records)
        rejected_rows = sum("non-numeric or empty population value" in item for item in rejections)
        input_rows += self._overview_input_count(options)
        if len({record.record_id for record in records}) != len(records):
            raise ValidationError("Population forecast workbook contains duplicate record IDs")
        warnings = list(options.get("limitations", ()))
        warnings.extend(
            (
                f"Age-versus-total checks: {age_checks} checked; maximum absolute discrepancy "
                f"{max_age_discrepancy:.6f} persons.",
                f"Prognoseraum-versus-district checks: {pgr_district_violations} differences "
                f"over 1 person; maximum absolute discrepancy {max_pgr_district:.6f} persons.",
                f"District-versus-Berlin checks: {district_berlin_violations} differences "
                f"over 1 person; maximum absolute discrepancy {max_district_berlin:.6f} persons.",
            )
        )
        return NormalizedBatch(
            records=tuple(records),
            input_rows=input_rows,
            rejected_rows=rejected_rows,
            rejections=tuple(rejections),
            warnings=tuple(warnings),
        )

    def _normalize_overview(self, path, raw, config, options, emit, rejections, records):
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        try:
            book = load_workbook(path, read_only=True, data_only=True)
        except (BadZipFile, OSError, ValueError) as exc:
            raise SchemaError("Scenario overview resource is not a readable XLSX workbook") from exc
        try:
            expected_sheets = options["expected_sheet_names"]["scenario_overview"]
            if book.sheetnames != expected_sheets:
                raise SchemaError(
                    f"Scenario overview workbook sheets changed: expected {expected_sheets}, "
                    f"found {book.sheetnames}"
                )
            name = options["overview_sheet_name"]
            if name not in book.sheetnames:
                raise SchemaError(f"Scenario overview sheet {name!r} is missing")
            sheet = book[name]
            rows = list(sheet.iter_rows(min_row=1, max_row=13, max_col=7, values_only=True))
            if rows[0][0] != options["expected_cells"]["overview_title"]:
                raise SchemaError(f"Overview title changed at {name}!A1")
            for scenario, column in _SCENARIO_COLUMNS.items():
                if rows[2][column - 1] != _SCENARIO_LABELS[scenario]:
                    raise SchemaError(
                        f"Overview scenario header changed at {name}!{get_column_letter(column)}3"
                    )
            expected_headers = options["expected_cells"]["overview_row4"]
            for column_text, expected in expected_headers.items():
                column = ord(column_text.upper()) - ord("A") + 1
                actual = rows[3][column - 1]
                if _normalize_header(actual) != _normalize_header(expected):
                    raise SchemaError(f"Overview header changed at {name}!{column_text}4")
            for row_number, expected in enumerate(options["overview_year_labels"], start=5):
                actual = rows[row_number - 1][0]
                if actual != expected:
                    raise SchemaError(
                        f"Overview year label changed at {name}!A{row_number}: {actual!r}"
                    )
            expected_years = {2024}
            expected_years.update(
                int(label[:-1])
                for label in options["overview_year_labels"]
                if isinstance(label, str) and re.fullmatch(r"\d{4}\*", label)
            )
            expected_matrix = {
                (scenario, year) for scenario in _SCENARIO_COLUMNS for year in expected_years
            }
            overview_matrix = set()
            for scenario, column in _SCENARIO_COLUMNS.items():
                for row_number in (10, 11, 12, 13):
                    label = rows[row_number - 1][0]
                    if row_number == 10:
                        year, status = 2024, "observed"
                    else:
                        if not isinstance(label, str) or not re.fullmatch(r"\d{4}\*", label):
                            raise SchemaError(
                                f"Overview forecast year must end with * at {name}!A{row_number}"
                            )
                        year, status = int(label[:-1]), "forecast"
                    value = rows[row_number - 1][column - 1]
                    value_cell = f"{get_column_letter(column)}{row_number}"
                    if (
                        isinstance(value, bool)
                        or not isinstance(value, (int, float))
                        or not math.isfinite(value)
                    ):
                        raise SchemaError(
                            f"Overview scenario value must be numeric at {name}!{value_cell}"
                        )
                    record = ForecastObservation(
                        source_id=config.dataset_id,
                        snapshot_id=raw.snapshot_id,
                        source_record_id=f"{name}!{get_column_letter(column)}{row_number}",
                        record_id=_record_id(
                            options["forecast_edition"],
                            scenario,
                            "land",
                            "11",
                            f"{year}-12-31",
                            "population_total",
                            None,
                            None,
                            "thousand_persons",
                        ),
                        area_code="11",
                        hierarchy="land",
                        boundary_edition=config.boundary_edition,
                        metric_code="population_total",
                        value=value,
                        unit="thousand_persons",
                        reference_period=f"{year}-12-31",
                        value_status=status,
                        scenario=scenario,
                        forecast_edition=options["forecast_edition"],
                        base_period=options["base_period"],
                        quality_flags=("rounded_to_thousands",),
                    )
                    records.append(record)
                    overview_matrix.add((scenario, year))
            if len(overview_matrix) != 12 or overview_matrix != expected_matrix:
                missing = sorted(expected_matrix - overview_matrix)
                raise SchemaError(f"Overview scenario/year matrix is incomplete; missing {missing}")
        finally:
            book.close()

    @staticmethod
    def _overview_input_count(options) -> int:
        return len(_SCENARIO_COLUMNS) * 4

    @staticmethod
    def _check_workbook_size(path: Path, options: dict) -> None:
        try:
            with ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > 1000 or sum(member.file_size for member in members) > int(
                    options.get("max_unpacked_bytes", 30_000_000)
                ):
                    raise SchemaError("Population forecast workbook exceeds unpacked size budget")
        except (BadZipFile, OSError) as exc:
            raise SchemaError("Population forecast resource is not a valid XLSX archive") from exc

    @staticmethod
    def _check_small_area_header(sheet_name, rows, options, config) -> None:
        cells = options["expected_cells"]
        title = rows[0][0] if rows and rows[0] else None
        scenario = rows[1][0] if len(rows) > 1 and rows[1] else None
        if title != cells["small_area_title"]:
            raise SchemaError(f"Small-area title changed at {sheet_name}!A1")
        if scenario != cells["small_area_scenario"]:
            raise SchemaError(
                f"Unexpected forecast scenario {scenario!r} at {sheet_name}!A2; "
                "all small-area sheets must use the middle variant"
            )
        base_value = rows[3][0]
        base_matches = (
            isinstance(base_value, (datetime, date))
            and base_value.date().isoformat() == config.parser_options["base_period"]
            if isinstance(base_value, datetime)
            else isinstance(base_value, date)
            and base_value.isoformat() == config.parser_options["base_period"]
        ) or (
            isinstance(base_value, str)
            and ("31.12.2024" in base_value or config.parser_options["base_period"] in base_value)
        )
        if not base_matches:
            raise SchemaError(f"Forecast base date changed at {sheet_name}!A4")
        if sheet_name == "00_Berlin":
            if rows[5][0] != cells["berlin_geography_label"]:
                raise SchemaError(f"Berlin geography label changed at {sheet_name}!A6")
        else:
            if rows[5][0] != cells["district_geography_label"]:
                raise SchemaError(f"District geography label changed at {sheet_name}!A6")
            district_code = sheet_name[:2]
            district_label = rows[5][2] if len(rows[5]) > 2 else None
            if not isinstance(district_label, str) or not re.fullmatch(
                rf"{district_code} - .+", district_label
            ):
                raise SchemaError(f"District code or name changed at {sheet_name}!C6")
            if isinstance(rows[3][2], (date, datetime)):
                date_value = rows[3][2].date() if isinstance(rows[3][2], datetime) else rows[3][2]
                if date_value.isoformat() != config.parser_options["base_period"]:
                    raise SchemaError(f"District base date changed at {sheet_name}!C4")
        years = options["target_years"]
        for start in (3,):
            if tuple(rows[7][start - 1 : start - 1 + len(years) + 1]) != tuple(range(2024, 2041)):
                raise SchemaError(f"Forecast year headers changed in {sheet_name} block at row 8")
        if sheet_name != "00_Berlin":
            for column, label in enumerate(rows[6], start=1):
                if label is not None and column >= 21:
                    if not isinstance(label, str) or not re.fullmatch(r"\d{4} - .+", label):
                        raise SchemaError(f"Invalid Prognoseraum label at {sheet_name}!{column}7")
                    start = column
                    if tuple(rows[7][start - 1 : start - 1 + 17]) != tuple(range(2024, 2041)):
                        raise SchemaError(
                            f"Forecast year headers changed in {sheet_name} block at row 8"
                        )
        for row_number, expected in enumerate(_AGE_LABELS, start=9):
            actual = rows[row_number - 1][0] if len(rows[row_number - 1]) else None
            if actual != expected:
                raise SchemaError(f"Age label changed at {sheet_name}!A{row_number}")
        if rows[108][0] != "Gesamt":
            raise SchemaError(f"Population total label changed at {sheet_name}!A109")

    def validate(self, batch: NormalizedBatch) -> QualityReport:
        errors = list(batch.rejections)
        if not batch.records:
            errors.append("Population forecast import contains no usable records")
        if any(not isinstance(record, ForecastObservation) for record in batch.records):
            errors.append("Population forecast batch contains a non-forecast record")
        if len({record.record_id for record in batch.records}) != len(batch.records):
            errors.append("Population forecast batch contains duplicate record IDs")
        if batch.rejected_rows:
            errors.append(f"{batch.rejected_rows} forecast data cells were rejected")
        return QualityReport(
            valid=not errors,
            input_rows=batch.input_rows,
            output_rows=len(batch.records),
            rejected_rows=batch.rejected_rows,
            warnings=batch.warnings,
            errors=tuple(errors),
        )

    def check_parent(self, config: DatasetConfig, records, parent_records) -> dict:
        parent_areas = [
            area
            for area in parent_records
            if isinstance(area, Area)
            and area.hierarchy == "planungsraum"
            and area.boundary_edition == config.boundary_edition
        ]
        lor_pgr_codes = {area.area_code[:4] for area in parent_areas}
        lor_district_codes = {area.area_code[:2] for area in parent_areas}
        forecast_pgr_codes = {
            record.area_code
            for record in records
            if isinstance(record, ForecastObservation) and record.hierarchy == "prognoseraum"
        }
        forecast_district_codes = {
            record.area_code
            for record in records
            if isinstance(record, ForecastObservation) and record.hierarchy == "bezirk"
        }
        missing = len(lor_pgr_codes - forecast_pgr_codes) + len(
            lor_district_codes - forecast_district_codes
        )
        unmatched = len(forecast_pgr_codes - lor_pgr_codes) + len(
            forecast_district_codes - lor_district_codes
        )
        if missing or unmatched or not parent_areas:
            raise IncompatibleSnapshotError(
                f"Forecast/LOR code coverage differs: {missing} missing, {unmatched} unmatched "
                "Prognoseraum/Bezirk codes"
            )
        if not forecast_pgr_codes or not forecast_district_codes:
            raise IncompatibleSnapshotError(
                "Forecast/LOR code coverage differs: missing or unmatched hierarchy codes"
            )
        return {
            "prognoseraum_count": len(forecast_pgr_codes),
            "bezirk_count": len(forecast_district_codes),
            "planungsraum_count": len(parent_areas),
            "join_method": ("Prognoseraum/Bezirk codes equal LOR 2021 Planungsraum code prefixes"),
            "resolution": "prognoseraum",
            "limitations": list(config.parser_options.get("limitations", ())),
        }
