"""Berlin household tables joined to the reviewed LOR code directory."""

from __future__ import annotations

import math
import re
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from app.connectors.base import FileConnector
from app.connectors.population import PopulationConnector, _label
from app.domain import (
    Area,
    AreaObservation,
    ConfigError,
    DatasetConfig,
    IncompatibleSnapshotError,
    NormalizedBatch,
    QualityReport,
    RawBundle,
    SchemaError,
)


class HouseholdsConnector(FileConnector):
    adapter_version = "1.0.0"

    def normalize(self, raw: RawBundle, config: DatasetConfig) -> NormalizedBatch:
        if len(raw.assets) != 1:
            raise SchemaError("Household import requires exactly one reviewed XLSX asset")
        options = config.parser_options
        if not config.reference_period:
            raise ConfigError("Household reference_period is required")
        if options.get("format") != "xlsx":
            raise ConfigError("Household format must be xlsx")
        if not isinstance(options.get("area_code_column"), str) or not re.fullmatch(
            r"[A-Z]+", options["area_code_column"]
        ):
            raise ConfigError("Household area_code_column must be an Excel column")
        unpacked_budget = options.get("max_unpacked_bytes", 30_000_000)
        if type(unpacked_budget) is not int or unpacked_budget <= 0:
            raise ConfigError("Household max_unpacked_bytes must be a positive integer")
        tables, _ = self._check_config(options)
        path = Path(raw.assets[0].path)
        with path.open("rb") as handle:
            prefix = handle.read(512).lstrip().lower()
        if prefix.startswith((b"<!doctype html", b"<html")) or b"<html" in prefix:
            raise SchemaError("Household resource returned HTML instead of a workbook")

        from openpyxl import load_workbook
        from openpyxl.utils import column_index_from_string

        try:
            with ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > unpacked_budget:
                    raise SchemaError("Household workbook exceeds unpacked size budget")
            book = load_workbook(path, read_only=True, data_only=False)
        except SchemaError:
            raise
        except (BadZipFile, OSError, ValueError) as exc:
            raise SchemaError("Household resource is not a readable XLSX workbook") from exc

        try:
            directory = options.get("code_directory", {})
            directory_sheet = directory.get("worksheet")
            if directory_sheet not in book.sheetnames:
                raise SchemaError(f"Household worksheet {directory_sheet!r} is missing")
            code_sheet = book[directory_sheet]
            self._check_cells(
                code_sheet,
                directory_sheet,
                directory.get("expected_cells", {}),
                require_checks=True,
            )
            code_column = directory.get("code_column")
            first_directory_row = directory.get("first_data_row")
            if not isinstance(code_column, str) or type(first_directory_row) is not int:
                raise ConfigError("Household code directory needs code_column and first_data_row")
            code_index = column_index_from_string(code_column)
            directory_codes: set[str] = set()
            input_rows = 0
            for row_number, values in enumerate(
                code_sheet.iter_rows(
                    min_row=first_directory_row, max_col=code_index, values_only=True
                ),
                first_directory_row,
            ):
                value = values[code_index - 1]
                if value is None:
                    continue
                input_rows += 1
                if not isinstance(value, str) or not re.fullmatch(r"[0-9]{8}", value.strip()):
                    raise SchemaError(
                        f"Household code directory {directory_sheet}!{code_column}{row_number} "
                        f"has invalid area code {value!r}"
                    )
                code = value.strip()
                if code in directory_codes:
                    raise SchemaError(f"Duplicate code directory area code {code}")
                directory_codes.add(code)

            if not directory_codes:
                raise SchemaError("Household code directory contains no area codes")

            area_rows: dict[str, dict[str, dict]] = {}
            total_rows: dict[str, dict[str, tuple[object, str, tuple[str, ...]]]] = {}
            warnings = list(options.get("limitations", ()))
            warnings.append(f"Code directory: {code_sheet['A2'].value}")
            rejected_areas: set[str] = set()
            structural_rejections: list[str] = []
            area_rejections: list[str] = []
            table_codes: dict[str, set[str]] = {}
            cell_integer: dict[tuple[str, str], bool] = {}
            for table in tables:
                sheet_name = table["worksheet"]
                for metric in table["metrics"]:
                    cell_integer[(sheet_name, metric["column"])] = metric.get("integer", True)
            for check in options.get("ratio_checks", []):
                sheet_name, column = self._split_ref(check["value"])
                cell_integer[(sheet_name, column)] = False

            for table in tables:
                sheet_name = table["worksheet"]
                if sheet_name not in book.sheetnames:
                    raise SchemaError(f"Household worksheet {sheet_name!r} is missing")
                sheet = book[sheet_name]
                self._check_cells(
                    sheet,
                    sheet_name,
                    table.get("expected_cells", {}),
                    require_checks=True,
                )
                check_columns = table["check_columns"]
                all_columns = set(check_columns) | {options["area_code_column"]}
                indexes = {column: column_index_from_string(column) - 1 for column in all_columns}
                area_column = options["area_code_column"]
                area_codes: set[str] = set()
                total_values = None
                total_row_count = 0
                first_data_row = table["first_data_row"]
                for row_number, values in enumerate(
                    sheet.iter_rows(
                        min_row=first_data_row,
                        max_col=max(indexes.values()) + 1,
                        values_only=True,
                    ),
                    first_data_row,
                ):
                    input_rows += 1
                    row = {column: values[index] for column, index in indexes.items()}
                    raw_code = row[area_column]
                    if raw_code is None:
                        district_header = (
                            isinstance(row.get("B"), str)
                            and bool(row["B"].strip())
                            and not any(
                                self._has_value(row[column])
                                for column in check_columns
                                if column != "B"
                            )
                        )
                        if not district_header and any(
                            self._has_value(row[column]) for column in check_columns
                        ):
                            area_rejections.append(
                                f"{sheet_name}!A{row_number}: values without area code"
                            )
                            rejected_areas.add(f"{sheet_name}!A{row_number}")
                        continue

                    label = raw_code.strip() if isinstance(raw_code, str) else None
                    if label in options.get("skip_labels", ()) or label == options.get(
                        "total_label"
                    ):
                        if label == options.get("total_label"):
                            total_row_count += 1
                            if total_row_count > 1:
                                structural_rejections.append(
                                    f"{sheet_name}: duplicate Berlin total row"
                                )
                                continue
                            total_values = self._parse_check_values(
                                row,
                                table,
                                options,
                                cell_integer,
                                None,
                                row_number,
                                structural_rejections,
                                rejected_areas,
                                is_total=True,
                            )
                        continue

                    if not isinstance(raw_code, str) or not re.fullmatch(
                        r"[0-9]{8}", raw_code.strip()
                    ):
                        area_rejections.append(
                            f"{sheet_name}!{area_column}{row_number}: invalid area code "
                            f"{raw_code!r}"
                        )
                        rejected_areas.add(f"{sheet_name}!A{row_number}")
                        continue

                    code = raw_code.strip()
                    if code in area_codes:
                        area_rejections.append(
                            f"{sheet_name}!A{row_number}: duplicate area code {code}"
                        )
                        rejected_areas.add(code)
                        continue
                    area_codes.add(code)
                    parsed = self._parse_check_values(
                        row,
                        table,
                        options,
                        cell_integer,
                        code,
                        row_number,
                        area_rejections,
                        rejected_areas,
                    )
                    area_rows.setdefault(code, {})[sheet_name] = {
                        "row": row_number,
                        "cells": parsed,
                    }

                table_codes[sheet_name] = area_codes
                expected_count = options.get("expected_area_count")
                if expected_count is not None and len(area_codes) != expected_count:
                    structural_rejections.append(
                        f"{sheet_name}: expected {expected_count} area rows, found {len(area_codes)}"
                    )
                missing = directory_codes - area_codes
                extra = area_codes - directory_codes
                if missing or extra:
                    for code in missing | extra:
                        rejected_areas.add(code)
                    structural_rejections.append(
                        f"{sheet_name}: code directory mismatch; missing {len(missing)} "
                        f"{sorted(missing)[:5]}, extra {len(extra)} {sorted(extra)[:5]}"
                    )
                if total_row_count == 0:
                    structural_rejections.append(f"{sheet_name}: Berlin total row is missing")
                elif total_values is not None:
                    total_rows[sheet_name] = total_values

            # Every configured table must contain the same complete area set.
            all_codes = set().union(*table_codes.values()) if table_codes else set()
            zero_denominator_refs: dict[str, set[str]] = {}
            for code in all_codes & directory_codes:
                cells = {}
                for table in tables:
                    table_data = area_rows.get(code, {}).get(table["worksheet"])
                    if table_data:
                        cells.update(
                            {
                                f"{table['worksheet']}!{column}": value
                                for column, value in table_data["cells"].items()
                            }
                        )
                for check in options.get("sum_checks", []):
                    refs = [check[0], *check[1]]
                    if all(cells.get(ref, (None, "missing", ()))[1] == "observed" for ref in refs):
                        lhs = cells[check[0]][0]
                        rhs = sum(cells[ref][0] for ref in check[1])
                        if lhs != rhs:
                            area_rejections.append(
                                f"{code}: sum check {check[0]} = {' + '.join(check[1])} failed "
                                f"({lhs} != {rhs})"
                            )
                            rejected_areas.add(code)
                for check in options.get("ratio_checks", []):
                    value = cells.get(check["value"], (None, "missing", ()))
                    numerator = cells.get(check["numerator"], (None, "missing", ()))
                    denominator = cells.get(check["denominator"], (None, "missing", ()))
                    if denominator[1] != "observed":
                        continue
                    if denominator[0] == 0:
                        if (
                            value[1] == "observed"
                            and value[0] != 0
                            or numerator[1] == "observed"
                            and numerator[0] != 0
                        ):
                            area_rejections.append(
                                f"{code}: ratio check {check['value']} has zero denominator "
                                f"{check['denominator']} but value {value[0]} "
                                f"(numerator {numerator[0]})"
                            )
                            rejected_areas.add(code)
                        else:
                            zero_denominator_refs.setdefault(code, set()).add(check["value"])
                        continue
                    if value[1] != "observed" or numerator[1] != "observed":
                        continue
                    expected = check["scale"] * numerator[0] / denominator[0]
                    if abs(value[0] - expected) > check["tolerance"]:
                        area_rejections.append(
                            f"{code}: ratio check {check['value']} = "
                            f"{check['scale']} × {check['numerator']} / "
                            f"{check['denominator']} failed ({value[0]} vs {expected})"
                        )
                        rejected_areas.add(code)

            for table in tables:
                sheet_name = table["worksheet"]
                total = total_rows.get(sheet_name)
                for metric in table["metrics"]:
                    column = metric["column"]
                    if not metric.get("integer", True):
                        continue
                    if total is None:
                        continue
                    observations = [
                        area_rows.get(code, {})
                        .get(sheet_name, {})
                        .get("cells", {})
                        .get(column, (None, "missing", ()))
                        for code in directory_codes
                    ]
                    if not all(status == "observed" for _, status, _ in observations):
                        warnings.append(
                            f"{sheet_name}!{column}: total check skipped; not all area values are observed"
                        )
                        continue
                    total_value, total_status, _ = total.get(column, (None, "missing", ()))
                    if total_status != "observed":
                        structural_rejections.append(
                            f"{sheet_name}!{column}: Berlin total is not observed"
                        )
                        continue
                    area_sum = sum(value for value, _, _ in observations)
                    if area_sum != total_value:
                        structural_rejections.append(
                            f"{sheet_name}!{column}: PLR sum {area_sum} differs from Berlin total "
                            f"{total_value}"
                        )

            records: list[AreaObservation] = []
            for code in sorted(directory_codes):
                if code in rejected_areas:
                    continue
                area_tables = area_rows.get(code, {})
                observations = []
                for table in tables:
                    sheet_name = table["worksheet"]
                    row_info = area_tables[sheet_name]
                    row_number = row_info["row"]
                    for metric in table["metrics"]:
                        value, status, flags = row_info["cells"][metric["column"]]
                        extra_flags = list(flags)
                        if metric.get("denominator_metric_code"):
                            extra_flags.extend(
                                (
                                    f"denominator:{metric['denominator_metric_code']}",
                                    "source_published_ratio",
                                )
                            )
                        if f"{sheet_name}!{metric['column']}" in zero_denominator_refs.get(
                            code, set()
                        ):
                            value, status = None, "not_applicable"
                            extra_flags.append("denominator_zero")
                        lower = metric.get("age_min")
                        upper = metric.get("age_max_exclusive")
                        record_id = (
                            f"{code}:{config.reference_period}:{metric['metric_code']}:"
                            f"{lower}:{upper}"
                        )
                        observations.append(
                            AreaObservation(
                                source_id=config.dataset_id,
                                snapshot_id=raw.snapshot_id,
                                source_record_id=f"{sheet_name}!{metric['column']}{row_number}",
                                record_id=record_id,
                                area_code=code,
                                hierarchy=options.get("hierarchy", "planungsraum"),
                                boundary_edition=config.boundary_edition,
                                metric_code=metric["metric_code"],
                                value=value,
                                unit=metric["unit"],
                                reference_period=config.reference_period,
                                value_status=status,
                                age_min=lower,
                                age_max_exclusive=upper,
                                quality_flags=tuple(extra_flags),
                            )
                        )

                for derived in options.get("derived_ratios", []):
                    numerator_ref = derived["numerator"]
                    denominator_ref = derived["denominator"]
                    cells = {}
                    for table in tables:
                        table_data = area_tables.get(table["worksheet"])
                        if table_data:
                            cells.update(
                                {
                                    f"{table['worksheet']}!{column}": value
                                    for column, value in table_data["cells"].items()
                                }
                            )
                    numerator = cells.get(numerator_ref, (None, "missing", ()))
                    denominator = cells.get(denominator_ref, (None, "missing", ()))
                    numerator_sheet, _ = self._split_ref(numerator_ref)
                    denominator_sheet, _ = self._split_ref(denominator_ref)
                    numerator_row = area_tables[numerator_sheet]["row"]
                    denominator_row = area_tables[denominator_sheet]["row"]
                    flags = [
                        "derived_ratio",
                        f"numerator:{derived['numerator_metric_code']}",
                        f"denominator:{derived['denominator_metric_code']}",
                    ]
                    if numerator[1] != "observed" or denominator[1] != "observed":
                        value, status = None, "missing"
                        flags.append("input_not_observed")
                    elif denominator[0] == 0:
                        value, status = None, "not_applicable"
                        flags.append("denominator_zero")
                    else:
                        value = round(
                            derived["scale"] * numerator[0] / denominator[0],
                            derived["round"],
                        )
                        status = "observed"
                    record_id = (
                        f"{code}:{config.reference_period}:{derived['metric_code']}:None:None"
                    )
                    observations.append(
                        AreaObservation(
                            source_id=config.dataset_id,
                            snapshot_id=raw.snapshot_id,
                            source_record_id=(
                                f"derived:{numerator_ref}{numerator_row}/"
                                f"{denominator_ref}{denominator_row}"
                            ),
                            record_id=record_id,
                            area_code=code,
                            hierarchy=options.get("hierarchy", "planungsraum"),
                            boundary_edition=config.boundary_edition,
                            metric_code=derived["metric_code"],
                            value=value,
                            unit=derived["unit"],
                            reference_period=config.reference_period,
                            value_status=status,
                            quality_flags=tuple(flags),
                        )
                    )
                records.extend(observations)

            rejections = tuple(structural_rejections + area_rejections)
            rejected_rows = len(rejected_areas) + len(structural_rejections)
            return NormalizedBatch(
                records=tuple(records),
                input_rows=max(input_rows, rejected_rows),
                rejected_rows=rejected_rows,
                rejections=rejections,
                warnings=tuple(warnings),
            )
        finally:
            book.close()

    @staticmethod
    def _check_config(options: dict) -> tuple[list[dict], dict]:
        tables = options.get("tables")
        if not isinstance(tables, list) or not tables:
            raise ConfigError("Household configuration requires tables")
        metrics_by_key = {}
        columns_by_table = {}
        for table in tables:
            sheet = table.get("worksheet")
            check_columns = table.get("check_columns")
            metrics = table.get("metrics")
            if (
                not isinstance(sheet, str)
                or not isinstance(check_columns, list)
                or not check_columns
                or any(
                    not isinstance(column, str) or not re.fullmatch(r"[A-Z]+", column)
                    for column in check_columns
                )
                or not isinstance(metrics, list)
            ):
                raise ConfigError(
                    "Every household table needs worksheet, check_columns, and metrics"
                )
            columns_by_table[sheet] = set(check_columns)
            for metric in metrics:
                if not isinstance(metric, dict) or not all(
                    isinstance(metric.get(key), str) and metric[key]
                    for key in ("column", "metric_code", "unit")
                ):
                    raise ConfigError("Every household metric needs column, metric_code, and unit")
                if not re.fullmatch(r"[A-Z]+", metric["column"]):
                    raise ConfigError("Household metric column must be an Excel column")
                if metric["column"] not in columns_by_table[sheet]:
                    raise ConfigError(
                        f"Household metric column {sheet}!{metric['column']} is not in check_columns"
                    )
                key = (
                    metric["metric_code"],
                    metric.get("age_min"),
                    metric.get("age_max_exclusive"),
                )
                if key in metrics_by_key:
                    raise ConfigError("Duplicate household metric mapping")
                metrics_by_key[key] = metric
        bands: dict[str, list[tuple[int, float]]] = {}
        derived_ratios = options.get("derived_ratios", [])
        if not isinstance(derived_ratios, list):
            raise ConfigError("Household derived_ratios must be a list")
        for derived in derived_ratios:
            if not isinstance(derived, dict) or not all(
                isinstance(derived.get(key), str) and derived[key]
                for key in ("metric_code", "unit")
            ):
                raise ConfigError("Every derived household ratio needs metric_code and unit")
        for metric in [metric for table in tables for metric in table["metrics"]] + derived_ratios:
            code = metric["metric_code"]
            lower, upper = metric.get("age_min"), metric.get("age_max_exclusive")
            if (
                lower is None
                and upper is not None
                or lower is not None
                and (
                    type(lower) is not int
                    or lower < 0
                    or upper is not None
                    and (type(upper) is not int or upper <= lower)
                )
            ):
                raise ConfigError("Invalid household age band bounds")
            if lower is not None:
                high = upper if upper is not None else math.inf
                existing = bands.setdefault(code, [])
                if any(lower < old_high and old_lower < high for old_lower, old_high in existing):
                    raise ConfigError(
                        "Overlapping household age bands are not permitted within a metric"
                    )
                existing.append((lower, high))

        for derived in derived_ratios:
            key = (derived["metric_code"], None, None)
            if key in metrics_by_key:
                raise ConfigError("Duplicate household metric mapping")
            metrics_by_key[key] = derived

        for check in options.get("sum_checks", []):
            if not isinstance(check, list) or len(check) != 2 or not isinstance(check[1], list):
                raise ConfigError(
                    "Household sum_checks must map a left reference to right references"
                )
            for reference in (check[0], *check[1]):
                HouseholdsConnector._validate_ref(reference, columns_by_table)
        for check in options.get("ratio_checks", []):
            for field in ("value", "numerator", "denominator"):
                HouseholdsConnector._validate_ref(check[field], columns_by_table)
        for derived in options.get("derived_ratios", []):
            for field in ("numerator", "denominator"):
                HouseholdsConnector._validate_ref(derived[field], columns_by_table)

        directory = options.get("code_directory", {})
        if not all(
            directory.get(key)
            for key in ("worksheet", "code_column", "first_data_row", "expected_cells")
        ):
            raise ConfigError("Household code_directory requires reviewed sheet and cell mappings")
        for table in tables:
            if type(table.get("first_data_row")) is not int or table["first_data_row"] < 1:
                raise ConfigError("Every household table needs first_data_row")
            if not table.get("expected_cells"):
                raise ConfigError("Every household table needs expected_cells")
        if type(directory["first_data_row"]) is not int or directory["first_data_row"] < 1:
            raise ConfigError("Household code directory first_data_row must be positive")
        if not re.fullmatch(r"[A-Z]+", directory["code_column"]):
            raise ConfigError("Household code directory code_column must be an Excel column")
        if options.get("expected_area_count") is not None and (
            type(options["expected_area_count"]) is not int or options["expected_area_count"] <= 0
        ):
            raise ConfigError("Household expected_area_count must be a positive integer")
        return tables, metrics_by_key

    @staticmethod
    def _validate_ref(reference: object, columns_by_table: dict[str, set[str]]) -> None:
        try:
            sheet, column = HouseholdsConnector._split_ref(reference)
        except (AttributeError, ValueError) as exc:
            raise ConfigError(f"Invalid household sheet-column reference: {reference!r}") from exc
        if sheet not in columns_by_table or column not in columns_by_table[sheet]:
            raise ConfigError(
                f"Household reference {reference!r} does not name a configured check column"
            )

    @staticmethod
    def _split_ref(reference: str) -> tuple[str, str]:
        sheet, column = reference.split("!", 1)
        if not sheet or not re.fullmatch(r"[A-Z]+", column):
            raise ValueError(reference)
        return sheet, column

    @staticmethod
    def _check_cells(sheet, sheet_name: str, expected: dict, require_checks: bool) -> None:
        if require_checks and not expected:
            raise ConfigError("Household workbook requires reviewed cell checks")
        for coordinate, value in expected.items():
            if _label(sheet[coordinate].value) != _label(value):
                raise SchemaError(f"Household XLSX schema changed at {sheet_name}!{coordinate}")

    @staticmethod
    def _has_value(value: object) -> bool:
        return isinstance(value, (int, float)) or (isinstance(value, str) and bool(value.strip()))

    @staticmethod
    def _parse_check_values(
        row: dict,
        table: dict,
        options: dict,
        cell_integer: dict[tuple[str, str], bool],
        code: str | None,
        row_number: int,
        rejections: list[str],
        rejected_areas: set[str],
        is_total: bool = False,
    ) -> dict[str, tuple[object, str, tuple[str, ...]]]:
        parsed = {}
        sheet_name = table["worksheet"]
        for column in table["check_columns"]:
            value = row[column]
            value_options = {
                "sentinels": options.get("sentinels", {}),
                "integer_counts": cell_integer.get((sheet_name, column), True),
                "thousands_separator": options.get("thousands_separator", ""),
                "decimal_separator": options.get("decimal_separator", "."),
            }
            try:
                parsed[column] = PopulationConnector._value(value, value_options)
            except ValueError as exc:
                prefix = "Berlin total" if is_total else f"{code}"
                rejections.append(f"{prefix} {sheet_name}!{column}{row_number}: {exc}")
                if code is not None:
                    rejected_areas.add(code)
                parsed[column] = (None, "missing", ())
        return parsed

    def validate(self, batch: NormalizedBatch) -> QualityReport:
        errors = list(batch.rejections)
        seen = set()
        if not batch.records:
            errors.append("Household import contains no usable observations")
        for record in batch.records:
            if not isinstance(record, AreaObservation):
                errors.append("Household batch contains a non-observation record")
                continue
            key = (
                record.area_code,
                record.hierarchy,
                record.boundary_edition,
                record.reference_period,
                record.metric_code,
                record.age_min,
                record.age_max_exclusive,
            )
            if key in seen:
                errors.append(f"Duplicate household observation key: {key}")
            seen.add(key)
        return QualityReport(
            valid=not errors,
            input_rows=batch.input_rows,
            output_rows=len(batch.records),
            rejected_rows=batch.rejected_rows,
            warnings=batch.warnings,
            errors=tuple(errors),
        )

    def check_parent(self, config: DatasetConfig, records, parent_records) -> dict:
        hierarchy = config.parser_options.get("hierarchy", "planungsraum")
        parent_areas = [record for record in parent_records if isinstance(record, Area)]
        observations = [record for record in records if isinstance(record, AreaObservation)]
        if any(area.boundary_edition != config.boundary_edition for area in parent_areas) or any(
            record.boundary_edition != config.boundary_edition for record in observations
        ):
            raise IncompatibleSnapshotError("Household and LOR boundary editions differ")
        parent_keys = {
            (area.hierarchy, area.area_code) for area in parent_areas if area.hierarchy == hierarchy
        }
        observed_keys = {
            (record.hierarchy, record.area_code)
            for record in observations
            if record.hierarchy == hierarchy
        }
        if parent_keys != observed_keys:
            missing = parent_keys - observed_keys
            extra = observed_keys - parent_keys
            raise IncompatibleSnapshotError(
                "Household/LOR coverage differs: "
                f"{len(missing)} missing areas {sorted(code for _, code in missing)[:5]}, "
                f"{len(extra)} unmatched areas {sorted(code for _, code in extra)[:5]}"
            )
        metric_keys = {
            (record.metric_code, record.age_min, record.age_max_exclusive)
            for record in observations
        }
        return {
            "area_count": len(observed_keys),
            "join_method": "exact hierarchy, boundary scheme, and complete area-code set",
            "metric_count": len(metric_keys),
            "reference_period": config.reference_period,
            "limitations": [
                "Code compatibility does not certify unchanged polygon boundaries between source editions."
            ],
        }
